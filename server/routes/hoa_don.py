"""Tải Hoá Đơn TCT + Quét hoá đơn Gmail trên trình duyệt (26.10.6).

Hai app của mục Kết nối ngoài, dùng chung một giao diện `/apps/tai-hoa-don/` và chung thư viện
`apps/tai-hoa-don/thd` (chép từ Tải Hoá Đơn TCT v1.27.0 của TrungKienHa). Mỗi nhóm API khoá
theo đúng app của nó:

- `tai_hoa_don`        doanh nghiệp, captcha + đăng nhập cổng thuế, tra cứu, tải XML / bản gốc,
                       xem hoá đơn, xuất Excel / bảng kê thuế / chi tiết mặt hàng / ZIP, cảnh báo
- `quet_hoa_don_gmail` cấu hình Gmail (mật khẩu ứng dụng), quét hộp thư, kết quả, tải file, Excel

Việc dài (tra cứu, tải, quét) chạy trong luồng nền thành "việc" có tiến độ và nhật ký; trình duyệt
hỏi `GET .../viec/{id}` để vẽ thanh tiến độ. Mật khẩu cổng thuế và mật khẩu ứng dụng Gmail mã
hoá bằng Fernet, khoá nằm ở `<STATE>/tai-hoa-don/.khoa` (quyền 600). Mọi tệp ở `<STATE>/tai-hoa-don/`.
"""
from __future__ import annotations

import datetime as dt
import importlib
import importlib.util
import io
import hashlib
import json
import os
import re
import sys
import threading
import time
import uuid
import zipfile
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, Request
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, Response

import giay_phep
import routes.ket_noi_ngoai as knn

GOC = "/apps/tai-hoa-don"
APP_DIR = Path(__file__).resolve().parents[2] / "apps" / "tai-hoa-don"
WEB_DIR = APP_DIR / "web"
WEB_FILES = {"index.html", "app.js", "invoice-ui.js", "styles.css"}
APP_TAI, APP_GMAIL, APP_CT = "tai_hoa_don", "quet_hoa_don_gmail", "bo_chung_tu_ke_toan"
LOAI = ("purchase", "sold")
THU_MUC_LOAI = {"purchase": "MuaVao", "sold": "BanRa"}


# ------------------------------------------------------------------ nạp thư viện

_THD = None


def thd(ten: str):
    """Module con của thư viện hoá đơn (nạp lười, dưới tên `minmin_thd` để không đụng module
    nào của server có cùng tên như `providers`, `analysis`)."""
    global _THD
    if _THD is None:
        goc = APP_DIR / "thd"
        spec = importlib.util.spec_from_file_location("minmin_thd", goc / "__init__.py",
                                                      submodule_search_locations=[str(goc)])
        m = importlib.util.module_from_spec(spec)
        sys.modules["minmin_thd"] = m
        spec.loader.exec_module(m)
        _THD = m
    return importlib.import_module("minmin_thd." + ten)


# ------------------------------------------------------------------ lưu trữ

def goc_du_lieu() -> Path:
    try:
        from config import STATE_DIR
    except Exception:  # pragma: no cover
        STATE_DIR = Path(os.environ.get("JAVIS_STATE_DIR") or Path.home() / ".minmin")
    p = Path(STATE_DIR) / "tai-hoa-don"
    p.mkdir(parents=True, exist_ok=True)
    return p


_KHOA_GHI = threading.RLock()


def _doc_json(p: Path, mac_dinh):
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return mac_dinh


def _ghi_json(p: Path, data) -> None:
    p.parent.mkdir(parents=True, exist_ok=True)
    tam = p.with_suffix(p.suffix + ".tmp")
    tam.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    os.replace(tam, p)


def _fernet():
    from cryptography.fernet import Fernet
    f = goc_du_lieu() / ".khoa"
    with _KHOA_GHI:
        if not f.exists():
            f.write_bytes(Fernet.generate_key())
            try:
                os.chmod(f, 0o600)
            except OSError:
                pass
    return Fernet(f.read_bytes().strip())


def ma_hoa(s: str) -> str:
    return _fernet().encrypt(str(s).encode("utf-8")).decode("ascii") if s else ""


def giai_ma(s: str) -> str:
    if not s:
        return ""
    try:
        return _fernet().decrypt(s.encode("ascii")).decode("utf-8")
    except Exception:
        return ""


def _cau_hinh_path() -> Path:
    return goc_du_lieu() / "cau-hinh.json"


def doc_cau_hinh() -> dict:
    d = _doc_json(_cau_hinh_path(), {})
    d.setdefault("doanh_nghiep", {})
    d.setdefault("gmail", {})
    d.setdefault("gmail_moc", {})
    return d


def ghi_cau_hinh(d: dict) -> None:
    with _KHOA_GHI:
        _ghi_json(_cau_hinh_path(), d)


def captcha_cong_khai() -> dict:
    c = doc_cau_hinh().get("captcha", {})
    return {"provider": c.get("provider", "2captcha"), "enabled": bool(c.get("enabled")),
            "keys": {p: bool(giai_ma(c.get(p, ""))) for p in ("2captcha", "omocaptcha")}}


def tao_solver(log=None):
    c = doc_cau_hinh().get("captcha", {})
    p = c.get("provider", "2captcha")
    key = giai_ma(c.get(p, "")) if p in ("2captcha", "omocaptcha") else ""
    if not c.get("enabled") or not key:
        return None
    return thd("captcha_solver").Solver(provider=p, api_key=key, log=log, local_first=False)


def ma_tra_cuu_hd(inv: dict) -> dict:
    """Không đoán mã từ số HĐ: chỉ dùng XML hoặc mã người dùng đã nhập."""
    xml = tuyet_doi(inv.get("_xml_path"))
    info = thd("providers").manual_download_info(xml.read_bytes(), inv) if xml else None
    info = info or {}
    source = inv.get("_gmail_lookup") or {}
    code = str(inv.get("_ma_tra_cuu") or info.get("code") or source.get("ma_tra_cuu") or "")
    if code.startswith("("):
        code = ""
    url = link_an_toan(info.get("url") or inv.get("_goc_link") or source.get("link"))
    note = ""
    if not xml and not code and not url:
        note = "Hoá đơn mới có trong danh sách tra cứu, chưa tải XML. Bấm Tải XML để nhận diện hãng và tìm mã tra cứu trong chứng từ."
    elif not code:
        note = "Nguồn đã đọc chưa có mã tra cứu. Có thể đối chiếu Gmail đã quét hoặc nhập mã từ email của nhà phát hành. Số hoá đơn không phải mã tra cứu."
    elif not url:
        note = "Đã tìm thấy mã tra cứu nhưng chưa xác định được địa chỉ trang của nhà phát hành. Bạn có thể sao chép mã; ứng dụng không mở một đường dẫn rỗng."
    return {"code": code, "url": url, "has_xml": bool(xml), "note": note,
            "provider": info.get("provider") or source.get("nha_cung_cap") or ""}


def chon_gmail(d: dict) -> list[dict]:
    kho = _doc_json(_gmail_kho_path(), {})
    if "ids" not in d:
        return sorted(kho.values(), key=lambda r: r.get("luc", 0), reverse=True)
    ids = d["ids"]
    if not isinstance(ids, list) or len(ids) > 10000 or any(not isinstance(i, str) for i in ids):
        raise ValueError("Danh sách hoá đơn không hợp lệ.")
    if not ids:
        raise ValueError("Chưa chọn hoá đơn nào.")
    if any(i not in kho for i in ids):
        raise ValueError("Danh sách đã thay đổi, tải lại trang rồi chọn lại.")
    return [kho[i] for i in dict.fromkeys(ids)]


def thong_tin_tra_cuu_gmail(r):
    """Restore old email tokens as actionable lookup links without rescanning mail."""
    from urllib.parse import quote, urlsplit
    scan = thd("gmail_scan")
    provider = r.get("nha_cung_cap") or ""
    info = scan.extract_link_info(provider, r.get("link") or "")
    code = r.get("ma_tra_cuu") or info.get("code") or ""
    info["code"] = code
    if provider == "bkav" and code:
        info["lookup_url"] = "https://tchd.ehoadon.vn/TCHD?MTC=" + quote(code, safe="")
    elif provider == "easyinvoice":
        token = info.get("token") or code
        if token:
            info["token"] = token
            host = urlsplit(r.get("link") or "").hostname or ""
            if not host.endswith(".easyinvoice.vn"):
                host = "tracuu.easyinvoice.vn"
            info["hosts"] = ["https://" + host]
            info["lookup_url"] = "https://" + host + "/Invoice/ViewFromEmail?token=" + quote(token, safe="")
    r["link"] = link_an_toan(info.get("lookup_url") or info.get("link") or r.get("link"))
    info["lookup_url"] = r["link"]
    return info


def tai_gmail_da_chon(v, records):
    prov = thd("providers")
    v.n = len(records)
    done = 0
    for r in records:
        if v.huy.is_set():
            break
        try:
            info = thong_tin_tra_cuu_gmail(r)
            v.log(f"[{v.i + 1}/{v.n}] {r.get('nha_cung_cap') or 'Hãng khác'} · HĐ {r.get('so') or r['id']}")
            if not tuyet_doi(r.get("pdf")):
                xml = tuyet_doi(r.get("xml"))
                inv = {"_bkav_mtc": r.get("ma_tra_cuu"), "_code": r.get("ma_tra_cuu"),
                       "_ma_tra_cuu": r.get("ma_tra_cuu")}
                solver = tao_solver(v.log)
                pdf = None
                xml_error = None
                if xml:
                    try:
                        _, pdf = prov.fetch_original_pdf(xml.read_bytes(), inv, solver=solver, log=v.log)
                    except Exception as e:
                        xml_error = e
                        v.log("Nguồn XML chưa tải được, thử link/mã từ email.")
                if not pdf:
                    if r.get("nha_cung_cap") == "bkav" and r.get("ma_tra_cuu"):
                        pdf = thd("providers.bkav").fetch_pdf("", inv, v.log, solver=solver)
                    else:
                        pdf = thd("gmail_scan").fetch_from_link(info, v.log)
                if not pdf:
                    raise ValueError(str(xml_error) if xml_error else "Thiếu link/mã hoặc hãng chưa hỗ trợ tải tự động. Bổ sung nguồn tra cứu.")
                if not pdf.lstrip().startswith(b"%PDF-"):
                    raise ValueError("Nguồn không trả PDF hợp lệ.")
                folder = goc_du_lieu() / "gmail" / "tai-lai"
                folder.mkdir(parents=True, exist_ok=True)
                path = folder / (uuid.uuid4().hex + ".pdf")
                path.write_bytes(pdf)
                r.update(pdf=_tuong_doi(str(path)), trang_thai="Đã tải PDF gốc", can_tai_tay=False)
            done += 1
            with _KHOA_GHI:
                kho = _doc_json(_gmail_kho_path(), {})
                # Giữ mã vừa sửa trên giao diện trong lúc việc nền chạy.
                current = kho.get(r["id"], {})
                current.update({k: r[k] for k in ("pdf", "trang_thai", "can_tai_tay", "link") if k in r})
                lien_ket_gmail(current)
                kho[r["id"]] = current
                _ghi_json(_gmail_kho_path(), kho)
        except Exception as e:
            error = str(e)[:220]
            v.log(f"Hoá đơn {r.get('so') or r['id']}: {error}")
            with _KHOA_GHI:
                kho = _doc_json(_gmail_kho_path(), {})
                if r["id"] in kho:
                    kho[r["id"]].update(trang_thai="Chưa tải được: " + error, can_tai_tay=True, link=r.get("link") or "")
                    _ghi_json(_gmail_kho_path(), kho)
        v.i += 1
    v.log(f"Xong: {done}/{v.n} hoá đơn có PDF; {v.n - done} chưa tải được.")
    return {"co_pdf": done, "thieu": v.n - done}


def chuan_mst(mst: str) -> str:
    m = re.sub(r"[^0-9\-]", "", str(mst or "")).strip("-")
    if not re.fullmatch(r"\d{10}(-\d{3})?", m):
        raise ValueError("Mã số thuế phải có 10 số (hoặc 10 số - 3 số chi nhánh).")
    return m


def _thu_muc_mst(mst: str) -> Path:
    return goc_du_lieu() / "du-lieu" / chuan_mst(mst)


def _kho_path(mst: str, loai: str) -> Path:
    return _thu_muc_mst(mst) / f"hoa-don-{loai}.json"


def doc_kho(mst: str, loai: str) -> dict:
    kho = _doc_json(_kho_path(mst, loai), {})
    # Older versions rendered HTTP-200 gateway errors as tax PDFs. Keep the
    # files for diagnosis, but never count, serve or export them as invoices.
    for inv in kho.values():
        if tuyet_doi(inv.get("_xml_path")):
            continue
        for field in ("_html_path", "_zip_path"):
            path = tuyet_doi(inv.get(field))
            if not path:
                continue
            try:
                with path.open("rb") as stream:
                    bad = thd("invoice_xml").is_error_document(stream.read(32768))
            except OSError:
                continue
            if bad:
                for key in ("_html_path", "_zip_path", "_pdf_path", "_pdf_render_version", "_html_assets"):
                    inv.pop(key, None)
                inv["_pdf_loi"] = "File cũ là trang báo lỗi cổng thuế, không phải chứng từ. Cần tải lại."
                break
    return kho


def ghi_kho(mst: str, loai: str, kho: dict) -> None:
    with _KHOA_GHI:
        _ghi_json(_kho_path(mst, loai), kho)


def khoa_hd(inv: dict) -> str:
    return thd("gdt_client").invoice_key(inv)


def ngay_hd(inv: dict) -> str:
    return str(inv.get("tdlap") or inv.get("ngay") or "")[:10]


def _tuong_doi(p: Optional[str]) -> Optional[str]:
    if not p:
        return None
    try:
        return str(Path(p).resolve().relative_to(goc_du_lieu().resolve()))
    except ValueError:
        return None


def tuyet_doi(p: Optional[str]) -> Optional[Path]:
    """Đường lưu trong kho là TƯƠNG ĐỐI gốc dữ liệu; ra ngoài gốc thì coi như không có."""
    if not p:
        return None
    goc = goc_du_lieu().resolve()
    q = (goc / p).resolve()
    try:
        q.relative_to(goc)
    except ValueError:
        return None
    return q if q.is_file() else None


# ------------------------------------------------------------------ việc nền

class Viec:
    def __init__(self, loai: str, mo_ta: str):
        self.id = uuid.uuid4().hex[:12]
        self.loai, self.mo_ta = loai, mo_ta
        self.trang_thai = "dang_chay"
        self.i = self.n = 0
        self.nhat_ky: list[str] = []
        self.ket_qua: dict = {}
        self.loi = ""
        self.huy = threading.Event()
        self.bat_dau = time.time()
        self.xong_luc = 0.0

    def log(self, s: str) -> None:
        self.nhat_ky.append(f"[{time.strftime('%H:%M:%S')}] {s}")
        if len(self.nhat_ky) > 400:
            del self.nhat_ky[:100]

    def ra_dict(self, tu: int = 0) -> dict:
        return {"id": self.id, "loai": self.loai, "mo_ta": self.mo_ta, "trang_thai": self.trang_thai,
                "i": self.i, "n": self.n, "loi": self.loi, "ket_qua": self.ket_qua,
                "nhat_ky": self.nhat_ky[tu:], "so_dong": len(self.nhat_ky),
                "bat_dau": self.bat_dau, "xong_luc": self.xong_luc}


VIEC: dict[str, Viec] = {}


def chay_viec(loai: str, mo_ta: str, ham) -> Viec:
    """Chạy `ham(viec)` trong luồng nền. Mỗi loại việc chỉ chạy một cái một lúc."""
    for v in VIEC.values():
        if v.loai == loai and v.trang_thai == "dang_chay":
            raise RuntimeError("Đang có một việc cùng loại chạy, đợi xong hoặc bấm Dừng.")
    v = Viec(loai, mo_ta)
    VIEC[v.id] = v
    for cu in sorted(VIEC.values(), key=lambda x: x.bat_dau)[:-30]:
        if cu.trang_thai != "dang_chay":
            VIEC.pop(cu.id, None)

    def _chay():
        try:
            v.ket_qua = ham(v) or {}
            v.trang_thai = "huy" if v.huy.is_set() else "xong"
        except Exception as e:
            if v.huy.is_set() or type(e).__name__ == "Cancelled":
                v.trang_thai = "huy"
                v.log("Đã dừng theo yêu cầu.")
            else:
                v.trang_thai, v.loi = "loi", str(e)
                v.log("Lỗi: " + str(e))
        finally:
            v.xong_luc = time.time()

    threading.Thread(target=_chay, name=f"hoa-don-{loai}", daemon=True).start()
    return v


# ------------------------------------------------------------------ cổng thuế

_CLIENT: dict[str, object] = {}


def _tct_base() -> Optional[str]:
    return os.environ.get("MINMIN_TCT_BASE") or None   # test trỏ vào cổng giả


def client_cho(mst: str):
    """Mỗi MST một phiên (giữ cookie tường lửa + token). Còn token đã lưu thì nạp lại."""
    mst = chuan_mst(mst)
    c = _CLIENT.get(mst)
    if c is None:
        c = thd("gdt_client").GDTClient(base=_tct_base())
        ph = _doc_json(goc_du_lieu() / "phien" / f"{mst}.json", {})
        tok = giai_ma(ph.get("token", ""))
        if tok:
            c.set_token(tok)
        _CLIENT[mst] = c
    return c


def luu_phien(mst: str, c) -> None:
    p = goc_du_lieu() / "phien" / f"{mst}.json"
    _ghi_json(p, {"token": ma_hoa(c.token or ""), "exp": c.token_exp()})
    try:
        os.chmod(p, 0o600)
    except OSError:
        pass


def da_dang_nhap(mst: str) -> tuple[bool, int]:
    try:
        c = client_cho(mst)
    except ValueError:
        return False, 0
    return c.token_valid(margin=120), c.token_exp()


def _can_dang_nhap(mst: str):
    c = client_cho(mst)
    if not c.token_valid(margin=60):
        raise RuntimeError("Phiên cổng thuế của MST này đã hết. Lấy captcha và đăng nhập lại.")
    return c


def _ngay(s: str) -> dt.date:
    try:
        return dt.date.fromisoformat(str(s)[:10])
    except ValueError:
        raise ValueError("Ngày không hợp lệ (dạng YYYY-MM-DD).")


def tra_cuu(v: Viec, mst: str, tu: dt.date, den: dt.date, loai_ds, he_ds) -> dict:
    c = _can_dang_nhap(mst)
    c.cancel_check = v.huy.is_set
    tong = {}
    v.n = len(loai_ds) * len(he_ds)
    for loai in loai_ds:
        kho = doc_kho(mst, loai)
        moi = 0
        for he in he_ds:
            v.log(f"Tra {'mua vào' if loai == 'purchase' else 'bán ra'}"
                  f"{' (máy tính tiền)' if he == 'sco-query' else ''} {tu:%d/%m/%Y} - {den:%d/%m/%Y}...")
            ds = c.query_invoices(tu, den, loai, he, progress=v.log)
            for inv in ds:
                k = khoa_hd(inv)
                cu = kho.get(k) or {}
                # Giữ lại phần MinMin tự thêm (đường file, chi tiết đã bóc) khi tra lại
                giu = {x: cu[x] for x in cu if x.startswith("_") and x not in ("_system", "_loai")}
                if not cu:
                    moi += 1
                kho[k] = {**inv, **giu}
            v.i += 1
        ghi_kho(mst, loai, kho)
        tong[loai] = moi
        v.log(f"Xong {'mua vào' if loai == 'purchase' else 'bán ra'}: thêm {moi} hoá đơn mới, kho có {len(kho)}.")
    if c.token_dead:
        raise RuntimeError("Phiên cổng thuế hết hạn giữa chừng, đăng nhập lại rồi tra tiếp.")
    return {"moi": tong}


def link_an_toan(u) -> str:
    """Link lấy từ XML / email của bên thứ ba: chỉ nhận http(s), chặn `javascript:` và các kiểu lạ."""
    u = str(u or "").strip()
    return u[:500] if re.match(r"^https?://[^\s\"'<>]+$", u, re.I) else ""


# ---- cài đặt tải (giống app PC: mẫu tên file, cách xếp thư mục, giữ zip, số luồng)

TOKEN_TEN_FILE = [("ngay_tai", "Ngày tải"), ("ngay_lap", "Ngày lập"), ("so_hd", "Số hoá đơn"),
                  ("ky_hieu", "Ký hiệu"), ("mst", "Mã số thuế"), ("ten_doitac", "Tên đối tác")]
TOKEN_TEN_EXCEL = [("mua_ban", "Mua/Bán"), ("tu_ngay", "Từ ngày"), ("den_ngay", "Đến ngày"),
                   ("mst", "MST"), ("ten_cty", "Tên công ty")]
XEP_THU_MUC = {"thang_lap": "Theo tháng lập hoá đơn", "ncc": "Theo tên đối tác (nhà cung cấp / khách)",
               "ngay_tai": "Theo ngày tải"}
CAI_DAT_MAC_DINH = {
    "ten_file": {"tokens": ["ngay_lap", "ten_doitac", "ky_hieu", "so_hd"], "keep_viet": False},
    "ten_excel": {"tokens": ["mua_ban", "tu_ngay", "den_ngay", "mst"], "keep_viet": False},
    "xep_thu_muc": "thang_lap", "giu_zip": False, "luong": 3,
}
LOAI_TEP = ("xml", "pdf", "goc", "html", "zip")
TEN_LOAI_TEP = {"xml": "XML thuế", "pdf": "PDF Trống đồng", "goc": "PDF gốc NCC", "html": "HTML Trống đồng",
                "zip": "ZIP gốc cổng thuế"}


def cai_dat_tai() -> dict:
    cd = json.loads(json.dumps(CAI_DAT_MAC_DINH))
    luu = doc_cau_hinh().get("tai") or {}
    for k in ("ten_file", "ten_excel"):
        if isinstance(luu.get(k), dict):
            hop = {t for t, _ in (TOKEN_TEN_FILE if k == "ten_file" else TOKEN_TEN_EXCEL)}
            toks = [t for t in luu[k].get("tokens") or [] if t in hop]
            if toks:
                cd[k]["tokens"] = toks
            cd[k]["keep_viet"] = bool(luu[k].get("keep_viet"))
    if luu.get("xep_thu_muc") in XEP_THU_MUC:
        cd["xep_thu_muc"] = luu["xep_thu_muc"]
    cd["giu_zip"] = bool(luu.get("giu_zip", cd["giu_zip"]))
    try:
        cd["luong"] = max(1, min(8, int(luu.get("luong", cd["luong"]))))
    except (TypeError, ValueError):
        pass
    return cd


def ten_file(inv: dict, cd: Optional[dict] = None) -> str:
    cd = cd or cai_dat_tai()
    return thd("gdt_client").build_filename(inv, tokens=cd["ten_file"]["tokens"], keep_viet=cd["ten_file"]["keep_viet"])


def ten_excel(mst: str, loai: str, tu: str, den: str, cd: Optional[dict] = None) -> str:
    cd = cd or cai_dat_tai()
    gdt = thd("gdt_client")
    ten_cty = (doc_cau_hinh()["doanh_nghiep"].get(mst) or {}).get("ten", "")
    gia_tri = {"mua_ban": {"purchase": "MuaVao", "sold": "BanRa"}.get(loai, "MuaBan"),
               "tu_ngay": tu.replace("-", ""), "den_ngay": den.replace("-", ""), "mst": mst,
               "ten_cty": gdt._san_name(ten_cty, 40)}
    ten = "_".join(v for v in (gia_tri.get(t, "") for t in cd["ten_excel"]["tokens"]) if v) or f"HoaDon_{mst}"
    if not cd["ten_excel"]["keep_viet"]:
        ten = gdt._strip_viet(ten)
    return gdt._san_name(ten, 150)


def _thu_muc_tai(mst: str, inv: dict, cd: Optional[dict] = None) -> Path:
    cd = cd or cai_dat_tai()
    gdt = thd("gdt_client")
    if cd["xep_thu_muc"] == "ncc":
        la = gdt._san_name(gdt.partner_name(inv), 60) or "KhongTenDoiTac"
        if not cd["ten_file"]["keep_viet"]:
            la = gdt._strip_viet(la)
    elif cd["xep_thu_muc"] == "ngay_tai":
        la = "Tai_" + dt.date.today().isoformat()
    else:
        la = ngay_hd(inv)[:7] or "khac"
    return _thu_muc_mst(mst) / THU_MUC_LOAI.get(inv.get("_loai"), "Khac") / la


def trinh_duyet_in_pdf() -> str:
    """Chrome/Chromium để in bản thể hiện ra PDF Trống đồng. Trên VPS thì là bản tải ở Công cụ."""
    try:
        import optional_tools
        p = optional_tools.duong_dan_chrome()
        if p:
            return p
    except Exception:
        pass
    return thd("exporter").find_browser() or ""


def tai_mot(c, mst: str, inv: dict, kieu, log, cd: Optional[dict] = None, browser: str = "") -> dict:
    """Tải các loại file được chọn cho MỘT hoá đơn, sửa `inv` tại chỗ.

    kieu: tập con của {"xml", "pdf", "goc"}. PDF Trống đồng và PDF gốc đều cần XML trước (bản
    thể hiện nằm trong zip cổng thuế; hãng phát hành nhận ra nhờ XML), nên XML luôn được tải."""
    cd = cd or cai_dat_tai()
    ixml, prov, ex = thd("invoice_xml"), thd("providers"), thd("exporter")
    base = ten_file(inv, cd)
    kq = {"xml": False, "pdf": False, "goc": False}
    xp = tuyet_doi(inv.get("_xml_path"))
    zp = tuyet_doi(inv.get("_zip_path"))
    can_zip = cd["giu_zip"] and not zp
    can_anh = "pdf" in kieu and inv.get("_html_path") and not inv.get("_html_assets")
    if not xp or can_zip or can_anh:
        folder = xp.parent if xp else _thu_muc_tai(mst, inv, cd)
        dung_zip_cu = bool(xp and zp)
        z = zp.read_bytes() if dung_zip_cu else c.download_xml_zip(inv)
        paths = ixml.save_invoice_files(str(folder), xp.stem if xp else base, z, keep_zip=cd["giu_zip"])
        for tr, k in (("_xml_path", "xml"), ("_html_path", "html"), ("_zip_path", "zip")):
            if paths.get(k):
                inv[tr] = _tuong_doi(paths[k])
        if paths.get("html"):
            inv["_html_assets"] = True
        if paths["xml"] and not xp:
            try:
                inv["_detail"] = ixml.parse_invoice(Path(paths["xml"]).read_bytes())
            except Exception as e:
                log(f"Không đọc được XML {base}: {e}")
            kq["xml"] = True
        elif not paths["xml"]:
            log(f"{base}: cổng thuế không trả XML, đã giữ file zip gốc.")
        if not dung_zip_cu:
            c.pause(c.delay_download)
        xp = tuyet_doi(inv.get("_xml_path"))
    # XML có thể đã tải ở lần trước, khi ấy vẫn cần đọc lại để tạo bản thể hiện.
    if xp and not inv.get("_detail"):
        try:
            inv["_detail"] = ixml.parse_invoice(xp.read_bytes())
        except Exception as e:
            log(f"Không đọc được XML {base}: {e}")
    if "pdf" in kieu and (not tuyet_doi(inv.get("_pdf_path")) or inv.get("_pdf_render_version") != 2):
        inv.pop("_pdf_path", None)  # file cũ vẫn ở đĩa; chỉ công bố bản in đã kiểm tra lại
        if not browser:
            inv["_pdf_loi"] = "Máy chủ chưa có trình duyệt để in PDF (cài ở Thiết lập > Công cụ)."
        else:
            nguon = tuyet_doi(inv.get("_html_path"))
            thu_muc = (xp.parent if xp else _thu_muc_tai(mst, inv, cd))
            stem = xp.stem if xp else base
            if not nguon and inv.get("_detail"):
                nguon = thu_muc / (stem + "_render.html")
                nguon.write_text(ixml.render_html(inv["_detail"]), encoding="utf-8")
            if nguon:
                try:
                    pdf = ex.html_to_pdf(str(nguon), str(thu_muc / (stem + ".pdf")), browser)
                    inv["_pdf_path"] = _tuong_doi(pdf)
                    inv["_pdf_render_version"] = 2
                    inv.pop("_pdf_loi", None)
                    kq["pdf"] = True
                except Exception as e:
                    inv["_pdf_loi"] = str(e)[:200]
                    log(f"In PDF Trống đồng lỗi ({base}): {e}")
            else:
                inv["_pdf_loi"] = "Chưa có HTML hoặc chi tiết XML hợp lệ để tạo PDF."
    if "goc" in kieu and xp and not tuyet_doi(inv.get("_goc_path")):
        try:
            pdf = None
            source = inv.get("_gmail_lookup")
            if isinstance(source, dict):
                try:
                    info = thong_tin_tra_cuu_gmail(dict(source))
                    if info.get("provider") == "myinvoice":
                        pdf = thd("providers.myinvoice").fetch_pdf(
                            xp.read_text(encoding="utf-8-sig"), {"_code": info.get("code")}, log)
                    else:
                        pdf = thd("gmail_scan").fetch_from_link(info, log)
                    ten = "Gmail"
                except Exception:
                    log("Link Gmail đã đối chiếu chưa trả PDF; thử nguồn XML và hãng phát hành.")
            if not pdf:
                ten, pdf = prov.fetch_original_pdf(xp.read_bytes(), inv, log=lambda s: log("   " + s), solver=tao_solver(log))
            if not pdf.lstrip().startswith(b"%PDF-"):
                raise prov.ProviderError("Nguồn không trả PDF hợp lệ.")
            g = xp.with_name(xp.stem + "_goc.pdf")
            g.write_bytes(pdf)
            inv["_goc_path"] = _tuong_doi(str(g))
            inv["_goc_nguon"] = ten
            inv.pop("_goc_loi", None)
            kq["goc"] = True
            log(f"Đã tải bản gốc ({ten}): {g.name}")
        except prov.ProviderError as e:
            info = prov.manual_download_info(xp.read_bytes(), inv) or {}
            url = info.get("url", "")
            inv["_goc_loi"] = str(e)[:200]
            if link_an_toan(url):
                inv["_goc_link"] = link_an_toan(url)
            log(f"Bản gốc chưa tải được: {inv['_goc_loi']}")
    thieu = [x for x in ("xml", "pdf", "goc")
             if (x == "xml" or x in kieu) and not tuyet_doi(inv.get(f"_{x}_path"))]
    if cd["giu_zip"] and not tuyet_doi(inv.get("_zip_path")):
        thieu.append("zip")
    kq["thieu"] = thieu
    if thieu:
        log(f"{base}: còn thiếu {', '.join(TEN_LOAI_TEP[x] for x in thieu)}.")
    return kq


def tai_ds(v: Viec, mst: str, loai: str, keys: list[str], kieu) -> dict:
    """Tải nhiều hoá đơn bằng vài luồng song song (mỗi luồng một phiên, chung token và chung mốc
    chờ khi cổng thuế báo 429), giống app PC."""
    from concurrent.futures import ThreadPoolExecutor
    c = _can_dang_nhap(mst)
    c.cancel_check = v.huy.is_set
    cd = cai_dat_tai()
    browser = trinh_duyet_in_pdf() if "pdf" in kieu else ""
    if "pdf" in kieu and not browser:
        v.log("Máy chủ chưa có trình duyệt nên chưa in được PDF Trống đồng. Vào Thiết lập > Công cụ, "
              "cài 'Trình duyệt' một lần (miễn phí). Vẫn tải XML.")
    kho = doc_kho(mst, loai)
    ds = [k for k in keys if k in kho]
    originals_before = sum(bool(tuyet_doi(kho[k].get("_goc_path"))) for k in ds)
    if "goc" in kieu and loai == "purchase":
        # One snapshot of all companies prevents a filtered selection hiding ambiguous matches.
        inventories = {m: doc_kho(m, "purchase") for m in doc_cau_hinh()["doanh_nghiep"]}
        inventories[mst] = kho
        linked = 0
        for ban in _doc_json(_gmail_kho_path(), {}).values():
            if v.huy.is_set():
                break
            try:
                hit = tim_hd_gmail(None, ban, inventories)
                if hit and hit[0] == mst and hit[1] in ds:
                    if ghep_chung_tu_gmail(kho[hit[1]], mst, ban):
                        linked += 1
                        v.log(f"Đã đối chiếu chứng từ Gmail: hoá đơn {kho[hit[1]].get('shdon')}.")
            except Exception as e:
                v.log(f"Không đối chiếu được một chứng từ Gmail: {type(e).__name__}.")
        v.log(f"Đối chiếu Gmail đã quét: bổ sung {linked} hoá đơn; chỉ ghép khi khớp duy nhất.")
    v.n = len(ds)
    n_luong = max(1, min(cd["luong"], len(ds) or 1))
    clients = [c] + [c.clone() for _ in range(n_luong - 1)]
    khoa = threading.Lock()
    phien_luong = threading.local()

    def gan_phien():
        with khoa:
            phien_luong.client = clients.pop()
    linked_originals = sum(bool(tuyet_doi(kho[k].get("_goc_path"))) for k in ds) - originals_before
    dem = {"ok": 0, "loi": 0, "goc": linked_originals, "pdf": 0, "het_phien": False}
    v.log(f"Tải {len(ds)} hoá đơn ({', '.join(TEN_LOAI_TEP[x] for x in kieu)}) bằng {n_luong} luồng.")

    def mot(i_k):
        i, k = i_k
        if v.huy.is_set() or dem["het_phien"]:
            return
        try:
            r = tai_mot(phien_luong.client, mst, kho[k], kieu, v.log, cd, browser)
            with khoa:
                dem["ok"] += int(not r["thieu"])
                dem["loi"] += int(bool(r["thieu"]))
                dem["goc"] += 1 if r["goc"] else 0
                dem["pdf"] += 1 if r["pdf"] else 0
        except Exception as e:
            ten = type(e).__name__
            with khoa:
                if ten == "TokenExpired":
                    dem["het_phien"] = True
                elif ten != "Cancelled":
                    dem["loi"] += 1
                    v.log(f"Lỗi {k}: {e}")
        with khoa:
            v.i += 1
            if v.i % 10 == 0:
                ghi_kho(mst, loai, kho)

    with ThreadPoolExecutor(max_workers=n_luong, initializer=gan_phien) as pool:
        list(pool.map(mot, enumerate(ds)))
    ghi_kho(mst, loai, kho)
    if dem["het_phien"]:
        raise RuntimeError("Phiên cổng thuế hết hạn giữa chừng. Đăng nhập lại rồi bấm Tải tiếp (file đã tải được giữ).")
    v.log(f"Xong: {dem['ok']} hoá đơn đủ file yêu cầu, {dem['pdf']} PDF Trống đồng mới, "
          f"{dem['goc']} bản gốc mới, {dem['loi']} hoá đơn lỗi hoặc thiếu file.")
    return {"ok": dem["ok"], "loi": dem["loi"], "goc": dem["goc"], "pdf": dem["pdf"]}


def loc_kho(mst: str, loai: str, tu: Optional[str], den: Optional[str]) -> list[dict]:
    a = (tu or "")[:10]
    b = (den or "")[:10]
    out = []
    for inv in doc_kho(mst, loai).values():
        d = ngay_hd(inv)
        if (a and d < a) or (b and d > b):
            continue
        out.append(inv)
    out.sort(key=lambda x: (ngay_hd(x), str(x.get("shdon"))), reverse=True)
    return out


def tom_tat(inv: dict) -> dict:
    gdt = thd("gdt_client")
    original = bool(tuyet_doi(inv.get("_goc_path")))
    xml = bool(tuyet_doi(inv.get("_xml_path")))
    original_status = "done" if original else "no-xml" if not xml else "failed" if inv.get("_goc_loi") else "pending"
    return {
        "key": khoa_hd(inv), "loai": inv.get("_loai"), "ngay": ngay_hd(inv),
        "so": inv.get("shdon"), "ky_hieu": f"{inv.get('khmshdon') or ''}{inv.get('khhdon') or ''}",
        "doi_tac": gdt.partner_name(inv),
        "mst_doi_tac": inv.get("nbmst") if inv.get("_loai") != "sold" else inv.get("nmmst"),
        "tien_hang": inv.get("tgtcthue") or 0, "thue": inv.get("tgtthue") or 0,
        "tong": inv.get("tgtttbso") or 0,
        "trang_thai": gdt.TTHAI.get(inv.get("tthai"), str(inv.get("tthai") or "")),
        "may_tinh_tien": inv.get("_system") == "sco-query", "tu_gmail": inv.get("_system") == "gmail",
        "co_xml": bool(tuyet_doi(inv.get("_xml_path"))),
        "co_html": bool(tuyet_doi(inv.get("_html_path"))) or bool(inv.get("_detail")),
        "co_goc": bool(tuyet_doi(inv.get("_goc_path"))), "goc_nguon": inv.get("_goc_nguon") or "",
        "co_trong_dong": bool(tuyet_doi(inv.get("_pdf_path"))),
        "co_pdf": bool(tuyet_doi(inv.get("_pdf_path"))), "pdf_loi": inv.get("_pdf_loi") or "",
        "co_zip": bool(tuyet_doi(inv.get("_zip_path"))),
        "goc_loi": inv.get("_goc_loi") or "", "goc_link": link_an_toan(inv.get("_goc_link")),
        "ma_tra_cuu": inv.get("_ma_tra_cuu") or "",
        "nha_cung_cap": thd("providers").issuer_name(inv),
        "can_tai_tay": not original,
        "goc_trang_thai": original_status,
    }


# ------------------------------------------------------------------ Gmail

def _gmail_kho_path() -> Path:
    return goc_du_lieu() / "gmail" / "ket-qua.json"


# ---- tài khoản Gmail: mật khẩu ứng dụng (IMAP) hoặc Đăng nhập bằng Google (OAuth, Gmail API)
#
# Tài khoản bật xác minh 2 bước kiểu g.co/sc, mã PIN, khoá bảo mật... không đăng nhập IMAP bằng
# mật khẩu thường được (Google trả "Application-specific password required"). Đăng nhập bằng
# Google đưa người dùng sang CHÍNH trang Google để qua mọi bước đó, MinMin chỉ nhận về quyền
# ĐỌC thư (gmail.readonly) và không bao giờ thấy mật khẩu.

GMAIL_SCOPE = "https://www.googleapis.com/auth/gmail.readonly"
GOOGLE_AUTH = "https://accounts.google.com/o/oauth2/v2/auth"
GOOGLE_TOKEN = "https://oauth2.googleapis.com/token"
GMAIL_API = "https://gmail.googleapis.com/gmail/v1/users/me"
_OAUTH_STATE: dict[str, tuple[float, str]] = {}


def ds_gmail(ch: Optional[dict] = None) -> dict:
    """{email: {"kieu": "mk"|"google", "mk"|"refresh": <mã hoá>}}. Đọc luôn cấu hình cũ (1 Gmail)."""
    ch = ch or doc_cau_hinh()
    ds = dict(ch.get("gmail_ds") or {})
    cu = ch.get("gmail") or {}
    if cu.get("email") and cu["email"] not in ds:
        ds[cu["email"]] = {"kieu": "mk", "mk": cu.get("mk", "")}
    return ds


def luu_gmail(email: str, tk: Optional[dict]) -> None:
    ch = doc_cau_hinh()
    ds = ds_gmail(ch)
    if tk is None:
        ds.pop(email, None)
    else:
        ds[email] = tk
    ch["gmail_ds"] = ds
    ch["gmail"] = {}
    ghi_cau_hinh(ch)


def google_client() -> tuple[str, str]:
    g = doc_cau_hinh().get("google_oauth") or {}
    return g.get("client_id", ""), giai_ma(g.get("client_secret", ""))


def google_access_token(refresh: str) -> str:
    import requests
    cid, sec = google_client()
    if not (cid and sec):
        raise RuntimeError("Chưa cài Google OAuth Client cho MinMin (mục Đăng nhập bằng Google).")
    r = requests.post(GOOGLE_TOKEN, data={"client_id": cid, "client_secret": sec, "refresh_token": refresh,
                                          "grant_type": "refresh_token"}, timeout=30)
    d = r.json() if r.content else {}
    if r.status_code != 200 or not d.get("access_token"):
        if d.get("error") == "invalid_grant":
            raise RuntimeError("Quyền đọc Gmail đã hết hạn hoặc bị thu hồi. Bấm Đăng nhập bằng Google lại cho tài khoản này.")
        raise RuntimeError(f"Google từ chối làm mới quyền: {d.get('error_description') or d.get('error') or r.status_code}")
    return d["access_token"]


def tao_bo_quet_google(email: str, refresh: str, log):
    """Bộ quét đọc thư qua Gmail API thay cho IMAP. Kế thừa GmailScanner nên phần nhận diện hoá
    đơn, tách đính kèm, cào PDF gốc giữ y nguyên; chỉ thay cách tìm và lấy thư."""
    import base64
    import requests
    gs = thd("gmail_scan")

    class BoQuetGoogle(gs.GmailScanner):
        def __init__(self):
            super().__init__(email, "", log=log)
            self._tok = ""
            self._ts_max = 0
            self._s = requests.Session()

        def connect(self):
            self.log(f"Kết nối Gmail {email} (Đăng nhập bằng Google)...")
            self._tok = google_access_token(refresh)
            self._s.headers["Authorization"] = "Bearer " + self._tok

        def close(self):
            self._s.close()

        def _get(self, url, params=None):
            for lan in (1, 2):
                r = self._s.get(url, params=params, timeout=60)
                if r.status_code == 401 and lan == 1:
                    self.connect()
                    continue
                if r.status_code == 429 and lan == 1:
                    time.sleep(5)
                    continue
                r.raise_for_status()
                return r.json()

        def _search_uids(self, d_from=None, d_to=None, since_uid=None):
            if since_uid:
                q = f"after:{int(since_uid)}"
            else:
                a = dt.datetime.strptime(d_from, "%d-%b-%Y").date()
                b = dt.datetime.strptime(d_to, "%d-%b-%Y").date()
                q = f"after:{a:%Y/%m/%d} before:{b:%Y/%m/%d}"
            ids, tok = [], None
            while True:
                p = {"q": q, "maxResults": 500, "includeSpamTrash": "false"}
                if tok:
                    p["pageToken"] = tok
                d = self._get(GMAIL_API + "/messages", p)
                ids += [m["id"] for m in d.get("messages") or []]
                tok = d.get("nextPageToken")
                if not tok or len(ids) >= 20000:
                    break
            return list(reversed(ids))   # API trả mới nhất trước; scan() cần cũ -> mới

        def _fetch_raw(self, mid):
            try:
                d = self._get(f"{GMAIL_API}/messages/{mid}", {"format": "raw"})
            except Exception as e:
                self.log(f"Không đọc được thư {mid}: {str(e)[:80]}")
                return None
            self._ts_max = max(self._ts_max, int(d.get("internalDate") or 0) // 1000)
            raw = d.get("raw") or ""
            return base64.urlsafe_b64decode(raw + "=" * (-len(raw) % 4))

        def scan(self, *a, **kw):
            truoc = int(kw.get("since_uid") or 0)
            out = super().scan(*a, **kw)
            self.last_uid = max(truoc, self._ts_max)   # mốc = thời điểm thư mới nhất đã đọc
            return out

    return BoQuetGoogle()


def _loi_gmail_de_hieu(e: Exception) -> str:
    s = str(e)
    if "Application-specific password required" in s or "application-specific" in s.lower():
        return ("Tài khoản này bật xác minh 2 bước nên Google không cho đăng nhập bằng mật khẩu Gmail thường. "
                "Bấm Đăng nhập bằng Google (qua được cả g.co/sc, mã PIN, khoá bảo mật), "
                "hoặc tạo Mật khẩu ứng dụng 16 ký tự rồi dán vào.")
    if "AUTHENTICATIONFAILED" in s or "Invalid credentials" in s:
        return ("Sai Gmail hoặc mật khẩu ứng dụng. Mật khẩu ứng dụng là 16 ký tự Google cấp riêng, "
                "không phải mật khẩu Gmail. Hoặc dùng Đăng nhập bằng Google.")
    return s


def _so_hd(x):
    t = str(x or "").strip()
    try:
        return int(t)
    except ValueError:
        return t.lstrip("0")


def tim_hd_gmail(det: Optional[dict], ban: dict, inventories=None) -> Optional[tuple]:
    """Hoá đơn mua vào nào trong kho Tải Hoá Đơn TCT ứng với một email hoá đơn.

    Khớp theo số hoá đơn + MST người bán + ký hiệu (cái nào email có). Chỉ nhận khi khớp ĐÚNG
    MỘT hoá đơn, để không bao giờ gắn nhầm bản gốc của hoá đơn này sang hoá đơn khác."""
    xml = tuyet_doi(ban.get("xml"))
    if det is None and xml:
        try:
            det = thd("invoice_xml").parse_invoice(xml.read_bytes())
        except Exception:
            return None
    inventories = inventories if inventories is not None else {
        m: doc_kho(m, "purchase") for m in doc_cau_hinh()["doanh_nghiep"]}
    pdf = tuyet_doi(ban.get("pdf"))
    if not det and pdf:
        identity = thd("invoice_identity")
        text = identity.read_pdf(pdf)
        numbers = identity.numbers(text)
        if len(numbers) != 1:
            return None
        hits = [(m, k) for m, rows in inventories.items() for k, inv in rows.items()
                if str(_so_hd(inv.get("shdon"))) in numbers and identity.matches(text, inv, m)]
        return hits[0] if len(hits) == 1 else None
    det = det or {}
    dsdn = list(inventories)
    mua = str(det.get("mua_mst") or "").strip()
    ung = [mua] if mua in dsdn else ([] if mua else dsdn)
    nb = str(det.get("ban_mst") or ban.get("ban_mst") or "").strip()
    kh = str(det.get("khhdon") or ban.get("ky_hieu") or "").upper()
    so = _so_hd(det.get("shdon") or ban.get("so"))
    if not so or not (nb or kh):
        return None
    trung = []
    for m in ung:
        for k, inv in inventories[m].items():
            if _so_hd(inv.get("shdon")) != so:
                continue
            if nb and str(inv.get("nbmst") or "") != nb:
                continue
            if kh:
                identity = thd("invoice_identity")
                source_serial = identity.serial({"khhdon": kh, "khmshdon": det.get("khmshdon")})
                if source_serial not in (str(inv.get("khhdon") or "").upper(), identity.serial(inv)):
                    continue
            trung.append((m, k))
    return trung[0] if len(trung) == 1 else None


def _chep_vao(inv: dict, mst: str, nguon: Path, duoi: str) -> Optional[str]:
    xp = tuyet_doi(inv.get("_xml_path"))
    thu_muc = xp.parent if xp else _thu_muc_tai(mst, inv)
    thu_muc.mkdir(parents=True, exist_ok=True)
    dich = thu_muc / ((xp.stem if xp else ten_file(inv)) + duoi)
    dich.write_bytes(nguon.read_bytes())
    return _tuong_doi(str(dich))


def ghep_chung_tu_gmail(inv, mst, ban, det=None):
    """Only called after unique identity reconciliation; never overwrite an existing original."""
    changed = False
    code = str(ban.get("ma_tra_cuu") or "")
    if code and not inv.get("_ma_tra_cuu"):
        inv.update(_ma_tra_cuu=code, _code=code)
        if ban.get("nha_cung_cap") == "bkav":
            inv["_bkav_mtc"] = code
        changed = True
    url = link_an_toan(ban.get("link"))
    if url and not inv.get("_goc_link"):
        inv["_goc_link"] = url
        changed = True
    source = {k: ban.get(k) or "" for k in ("nha_cung_cap", "ma_tra_cuu")}
    source["link"] = url
    if (url or code) and source != inv.get("_gmail_lookup"):
        inv["_gmail_lookup"] = source
        changed = True
    for src, dest, suffix in (("pdf", "_goc_path", "_goc.pdf"), ("xml", "_xml_path", ".xml")):
        path = tuyet_doi(ban.get(src))
        if path and not tuyet_doi(inv.get(dest)):
            if src == "pdf":
                with path.open("rb") as f:
                    if f.read(5) != b"%PDF-":
                        continue
            inv[dest] = _chep_vao(inv, mst, path, suffix)
            changed = True
            if src == "pdf":
                inv["_goc_nguon"] = "Gmail"
                inv.pop("_goc_loi", None)
            elif det:
                inv["_detail"] = det
    return changed


def gmail_trong_dong(ban: dict, cache: Optional[dict] = None) -> Optional[Path]:
    """Only expose the tax rendering of an invoice already reconciled with a managed company."""
    link = ban.get("lien_ket") or {}
    mst, key = link.get("mst"), link.get("key")
    if not mst or not key or mst not in doc_cau_hinh()["doanh_nghiep"]:
        return None
    if cache is None:
        cache = {}
    if mst not in cache:
        cache[mst] = doc_kho(mst, "purchase")
    return tuyet_doi((cache[mst].get(key) or {}).get("_pdf_path"))


def lien_ket_gmail(ban: dict, det: Optional[dict] = None) -> Optional[str]:
    """Ghép hoá đơn quét được từ Gmail vào kho Tải Hoá Đơn TCT. Trả "ghep" | "moi" | None.

    - Có sẵn trong kho (đã tra cứu cổng thuế): thêm PDF gốc (nếu chưa có) và XML (nếu chưa có).
    - Chưa có mà XML cho biết người mua là một doanh nghiệp đang quản lý: thêm mới vào mua vào.
    """
    pdf, xml = tuyet_doi(ban.get("pdf")), tuyet_doi(ban.get("xml"))
    if not (pdf or xml) and not (ban.get("ban_mst") and ban.get("ky_hieu") and ban.get("so")):
        return None
    if det is None and xml:
        try:
            det = thd("invoice_xml").parse_invoice(xml.read_bytes())
        except Exception:
            det = None
    with _KHOA_GHI:
        hit = tim_hd_gmail(det, ban)
        if hit:
            mst, key = hit
            kho = doc_kho(mst, "purchase")
            inv = kho[key]
            doi = ghep_chung_tu_gmail(inv, mst, ban, det)
            if doi:
                ghi_kho(mst, "purchase", kho)
            ban["lien_ket"] = {"mst": mst, "key": key}
            return "ghep" if doi else None
        mua = str((det or {}).get("mua_mst") or "").strip()
        if not det or not det.get("shdon") or not det.get("khhdon") or not det.get("ban_mst") or mua not in doc_cau_hinh()["doanh_nghiep"]:
            return None
        inv = {"nbmst": det.get("ban_mst"), "nbten": det.get("ban_ten"), "nmmst": mua, "nmten": det.get("mua_ten"),
               "khmshdon": det.get("khmshdon"), "khhdon": det.get("khhdon"), "shdon": _so_hd(det.get("shdon")),
               "tdlap": (str(det.get("ngay") or "")[:10] + "T00:00:00") if det.get("ngay") else "",
               "tgtcthue": det.get("tien_chua_thue") or 0, "tgtthue": det.get("tien_thue") or 0,
               "tgtttbso": det.get("tong_tien") or 0, "dvtte": det.get("dvtte") or "VND", "mccqt": det.get("mccqt"),
               "_loai": "purchase", "_system": "gmail", "_detail": det}
        key = khoa_hd(inv)
        kho = doc_kho(mua, "purchase")
        if key in kho:
            return None
        if xml:
            inv["_xml_path"] = _chep_vao(inv, mua, xml, ".xml")
        if pdf:
            inv["_goc_path"] = _chep_vao(inv, mua, pdf, "_goc.pdf")
            inv["_goc_nguon"] = "Gmail"
        kho[key] = inv
        ghi_kho(mua, "purchase", kho)
        ban["lien_ket"] = {"mst": mua, "key": key}
        return "moi"


def quet_gmail(v: Viec, tu: dt.date, den: dt.date, chi_moi: bool, email: str = "") -> dict:
    ixml = thd("invoice_xml")
    ds = ds_gmail()
    if not ds:
        raise RuntimeError("Chưa thêm tài khoản Gmail nào.")
    mail = email if email in ds else next(iter(ds))
    tk = ds[mail]
    khoa_moc = ("google:" if tk.get("kieu") == "google" else "") + mail
    moc = int(doc_cau_hinh()["gmail_moc"].get(khoa_moc, 0) or 0) if chi_moi else 0
    if tk.get("kieu") == "google":
        sc = tao_bo_quet_google(mail, giai_ma(tk.get("refresh", "")), v.log)
    else:
        mk = giai_ma(tk.get("mk", ""))
        if not mk:
            raise RuntimeError("Tài khoản này chưa có mật khẩu ứng dụng.")
        sc = thd("gmail_scan").GmailScanner(mail, mk, log=v.log)
    try:
        sc.connect()
    except Exception as e:
        raise RuntimeError(_loi_gmail_de_hieu(e))
    kho = _doc_json(_gmail_kho_path(), {})
    thu_muc = goc_du_lieu() / "gmail" / dt.date.today().strftime("%Y-%m")
    thu_muc.mkdir(parents=True, exist_ok=True)
    moi = 0
    ghep: dict = {}
    def luu_ban(r):
        nonlocal moi
        det = None
        if r.get("xml"):
            try:
                det = ixml.parse_invoice(r["xml"])
            except Exception:
                det = None
        so = (det or {}).get("shdon") or r.get("_so_hd") or ""
        nb = (det or {}).get("ban_mst") or r.get("_seller_mst") or ""
        serial = (det or {}).get("khhdon") or r.get("_khhdon") or ""
        if nb and serial and so:
            rid = re.sub(r"[^0-9A-Za-z_-]", "", f"{nb}_{serial}_{so}")
        else:
            # Incomplete identities must not collapse unrelated emails into '__'.
            source = r.get("_source_id") or repr((r.get("from"), r.get("subject"), r.get("date"), r.get("code")))
            # Portal PDFs may contain a fresh timestamp/document ID on every fetch.
            # Email attachments remain byte-keyed to keep multiple PDFs separate.
            portal = bool(r.get("lookup_url") and not r.get("xml"))
            payload = str(r["lookup_url"]).encode() if portal else (r.get("xml") or r.get("pdf") or b"")
            digest = hashlib.sha256((mail + source).encode() + payload).hexdigest()[:24]
            rid = "mail_" + digest
            if portal and r.get("subject") and r.get("from") and r.get("date"):
                # Reuse existing byte-keyed rows without deleting their saved files.
                legacy = [b for b in kho.values() if b.get("id", "").startswith("mail_")
                          and b.get("hop_thu") == mail and b.get("tieu_de") == r["subject"][:200]
                          and b.get("tu") == r["from"][:120] and b.get("ngay_mail") == r["date"]
                          and b.get("link") == link_an_toan(r["lookup_url"])]
                if legacy:
                    rid = min(legacy, key=lambda b: (b.get("luc", 0), b["id"]))["id"]
        cu = kho.get(rid) or {}
        bo_sung = any(r.get(t) and not tuyet_doi(cu.get(t)) for t in ("xml", "pdf"))
        if cu and not bo_sung:
            if not cu.get("ma_tra_cuu") and (r.get("_mtc") or r.get("code")):
                cu["ma_tra_cuu"] = r.get("_mtc") or r.get("code")
                cu["link"] = link_an_toan(r.get("lookup_url")) or cu.get("link")
                lien_ket_gmail(cu)
            if not kho[rid].get("lien_ket"):
                ket = lien_ket_gmail(kho[rid])
                if ket:
                    ghep[ket] = ghep.get(ket, 0) + 1
            return
        stem = re.sub(r"[^0-9A-Za-z_-]", "_", rid)[:80]
        ban = {"id": rid, "hop_thu": mail, "tieu_de": r.get("subject", "")[:200], "tu": r.get("from", "")[:120],
               "ngay_mail": r.get("date", ""), "nha_cung_cap": r.get("provider") or "",
               "trang_thai": r.get("status", ""), "can_tai_tay": bool(r.get("needs_manual")),
               "ma_tra_cuu": r.get("_mtc") or r.get("code") or "", "link": link_an_toan(r.get("lookup_url")),
               "ban_ten": (det or {}).get("ban_ten") or r.get("_seller") or "", "ban_mst": nb,
               "so": so, "ngay_hd": (det or {}).get("ngay") or "",
               "ky_hieu": ((det or {}).get("khhdon") or r.get("_khhdon") or "").upper(),
               "tong": (det or {}).get("tong_tien") or 0, "thue": (det or {}).get("tien_thue") or 0,
               "xml": cu.get("xml"), "pdf": cu.get("pdf"), "luc": time.time()}
        if r.get("xml"):
            p = thu_muc / f"{stem}.xml"
            p.write_bytes(r["xml"])
            ban["xml"] = _tuong_doi(str(p))
        if r.get("pdf"):
            p = thu_muc / f"{stem}.pdf"
            p.write_bytes(r["pdf"])
            ban["pdf"] = _tuong_doi(str(p))
        kho[rid] = ban
        moi += 1
        ket = lien_ket_gmail(ban, det)
        if ket:
            ghep[ket] = ghep.get(ket, 0) + 1
    def nhan_ban(r):
        with _KHOA_GHI:
            # Keep earlier invoices durable even if the scan or server stops.
            kho.clear()
            kho.update(_doc_json(_gmail_kho_path(), {}))
            luu_ban(r)
            _ghi_json(_gmail_kho_path(), kho)

    try:
        def tien_do(i, n):
            v.i, v.n = i, n
        recs = sc.scan(tu, den, want_download=True, progress=tien_do, since_uid=(moc or None),
                       cancel_check=v.huy.is_set, on_record=nhan_ban)
    finally:
        sc.close()
    if getattr(sc, "last_uid", 0) and not v.huy.is_set():
        ch = doc_cau_hinh()
        ch["gmail_moc"][khoa_moc] = int(sc.last_uid)
        ghi_cau_hinh(ch)
    for r in recs:
        nhan_ban(r)
    v.log(f"Lưu {moi} hoá đơn mới từ Gmail. Tổng trong kho: {len(kho)}.")
    if ghep:
        v.log(f"Đã ghép vào Tải Hoá Đơn TCT: {ghep.get('ghep', 0)} hoá đơn có sẵn được thêm bản gốc / XML, "
              f"{ghep.get('moi', 0)} hoá đơn mới thêm vào kho mua vào.")
    return {"moi": moi, "tong": len(kho), "ghep": ghep.get("ghep", 0), "them_vao_kho": ghep.get("moi", 0)}


def xuat_gmail_excel(ds: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "HoaDonGmail"
    tieu = ["Ngày HĐ", "Số HĐ", "Người bán", "MST người bán", "Tổng tiền", "Thuế", "Nhà cung cấp HĐ",
            "Trạng thái", "Cần tải tay", "Mã tra cứu", "Link tra cứu", "Tiêu đề email", "Người gửi"]
    ws.append(tieu)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in ds:
        ws.append([r.get("ngay_hd"), r.get("so"), r.get("ban_ten"), r.get("ban_mst"), r.get("tong") or 0,
                   r.get("thue") or 0, r.get("nha_cung_cap"), r.get("trang_thai"),
                   "Có" if r.get("can_tai_tay") else "", r.get("ma_tra_cuu"), r.get("link"),
                   r.get("tieu_de"), r.get("tu")])
    for col, w in zip("ABCDEFGHIJKLM", (12, 10, 40, 16, 14, 12, 14, 40, 10, 16, 40, 50, 30)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()



# ------------------------------------------------------------------ bộ chứng từ kế toán

def _html(s) -> str:
    import html
    return html.escape(str(s or ""), quote=True)


def _slug(s: str, limit: int = 72) -> str:
    raw = re.sub(r"[^0-9A-Za-zÀ-ỹĐđ._ -]+", "", str(s or "")).strip()
    raw = re.sub(r"\s+", " ", raw)[:limit].strip(" ._")
    return raw or "chung-tu"


def _so_tien(v) -> str:
    try:
        n = float(v or 0)
    except Exception:
        n = 0
    if abs(n - round(n)) < 0.01:
        return f"{int(round(n)):,}".replace(",", ".")
    return f"{n:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")


def _ngay_vn(s) -> str:
    s = str(s or "")[:10]
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        return f"{s[8:10]}/{s[5:7]}/{s[:4]}"
    return s


def _cty(mst: str) -> dict:
    d = doc_cau_hinh().get("doanh_nghiep", {}).get(mst, {})
    return {"mst": mst, "ten": d.get("ten") or f"Doanh nghiệp {mst}", "dia_chi": d.get("dia_chi", "")}


def _ben(inv: dict) -> dict:
    return {"ten": inv.get("nbten") or inv.get("doi_tac") or inv.get("ten") or "Đối tác",
            "mst": inv.get("nbmst") or inv.get("mst_doi_tac") or inv.get("mst") or "",
            "dia_chi": inv.get("nbdchi") or inv.get("dia_chi") or ""}


def _doc_frame(tieu_de: str, body: str, kieu: str = "ke_toan") -> str:
    mau = {"ke_toan": ("#0f766e", "#f8f4ec"), "trang_trong": ("#5b3fd6", "#f7f1e7"), "toi_gian": ("#111827", "#ffffff")}.get(kieu, ("#0f766e", "#f8f4ec"))
    return f'''<!doctype html><html lang="vi"><head><meta charset="utf-8"><title>{_html(tieu_de)}</title>
<style>@page{{size:A4;margin:20mm}}body{{font-family:"Times New Roman",serif;color:#111827;background:{mau[1]};font-size:13pt;line-height:1.45}}.page{{background:white;border:1px solid #e5ded1;border-radius:10px;padding:22px;box-shadow:0 8px 28px rgba(49,35,16,.08)}}h1{{text-align:center;color:{mau[0]};font-size:19pt;text-transform:uppercase;margin:8px 0 2px}}h2{{font-size:14pt;color:{mau[0]};margin:18px 0 8px}}.sub{{text-align:center;font-style:italic;margin-bottom:18px}}table{{width:100%;border-collapse:collapse;margin:10px 0}}td,th{{border:1px solid #7c7c7c;padding:6px;vertical-align:top}}th{{background:#f4efe6}}.no-border td{{border:0}}.right{{text-align:right}}.center{{text-align:center}}.muted{{color:#6b7280;font-size:11pt}}.sign td{{height:90px;text-align:center;border:0;font-weight:bold}}.line{{border-bottom:1px dotted #777;display:inline-block;min-width:180px}}</style></head><body><div class="page">{body}</div></body></html>'''


def _doc_bytes(tieu_de: str, body: str, kieu: str) -> bytes:
    return _doc_frame(tieu_de, body, kieu).encode("utf-8")


def _hoa_don_id(inv: dict) -> str:
    return f"{inv.get('khhdon') or inv.get('ky_hieu') or ''}-{inv.get('shdon') or inv.get('so') or ''}".strip("-") or khoa_hd(inv)


def _dong_tien(inv: dict) -> str:
    return f"Tiền hàng {_so_tien(inv.get('tien_hang') or inv.get('tgtcthue'))} đ; thuế {_so_tien(inv.get('thue') or inv.get('tgtthue'))} đ; tổng thanh toán {_so_tien(inv.get('tong') or inv.get('tgtttbso'))} đ"


def _body_chung_tu(tieu_de: str, cty: dict, ben: dict, inv: dict, loai: str, noi_dung: str) -> str:
    hd = _hoa_don_id(inv)
    return f'''<p class="center"><b>{_html(cty['ten'])}</b><br>MST: {_html(cty['mst'])}</p><h1>{_html(tieu_de)}</h1><p class="sub">Lập tự động từ hóa đơn {_html(hd)} ngày {_html(_ngay_vn(ngay_hd(inv)))}</p>
<table class="no-border"><tr><td><b>Đối tác:</b> {_html(ben['ten'])}</td><td><b>MST:</b> {_html(ben['mst'])}</td></tr><tr><td colspan="2"><b>Địa chỉ:</b> {_html(ben.get('dia_chi'))}</td></tr></table>
<table><thead><tr><th>Diễn giải</th><th>Số hóa đơn</th><th>Ngày</th><th class="right">Tiền hàng</th><th class="right">Thuế</th><th class="right">Thanh toán</th></tr></thead><tbody><tr><td>{_html(noi_dung)}</td><td>{_html(hd)}</td><td>{_html(_ngay_vn(ngay_hd(inv)))}</td><td class="right">{_so_tien(inv.get('tien_hang') or inv.get('tgtcthue'))}</td><td class="right">{_so_tien(inv.get('thue') or inv.get('tgtthue'))}</td><td class="right">{_so_tien(inv.get('tong') or inv.get('tgtttbso'))}</td></tr></tbody></table>
<p><b>Ghi chú kế toán:</b> {_html(_dong_tien(inv))}. Trạng thái PDF gốc NCC: {_html('đã có' if inv.get('_goc_path') else inv.get('_goc_loi') or 'chưa có')}.</p>
<table class="sign"><tr><td>Người lập<br><span class="muted">(ký, ghi rõ họ tên)</span></td><td>Kế toán trưởng<br><span class="muted">(ký, ghi rõ họ tên)</span></td><td>Giám đốc<br><span class="muted">(ký, đóng dấu)</span></td></tr></table>'''


def _invoice_docs(cty: dict, inv: dict, loai: str, nhom: set[str], kieu: str) -> list[tuple[str, bytes]]:
    ben = _ben(inv)
    docs = []
    if loai == "purchase" and "mua_vao" in nhom:
        items = [("01_Phieu_nhap_kho.doc", "Phiếu nhập kho", "Nhập hàng hóa/dịch vụ theo hóa đơn nhà cung cấp"), ("02_De_nghi_thanh_toan.doc", "Giấy đề nghị thanh toán", "Đề nghị thanh toán công nợ phát sinh theo hóa đơn"), ("03_Phieu_chi_Uy_nhiem_chi.doc", "Phiếu chi / Ủy nhiệm chi", "Thanh toán cho nhà cung cấp theo hóa đơn"), ("04_Bien_ban_giao_nhan.doc", "Biên bản giao nhận", "Xác nhận giao nhận hàng hóa/dịch vụ theo hóa đơn"), ("05_Doi_chieu_cong_no_phai_tra.doc", "Biên bản đối chiếu công nợ phải trả", "Đối chiếu công nợ phải trả với nhà cung cấp")]
    elif loai == "sold" and "ban_ra" in nhom:
        items = [("01_Phieu_xuat_kho.doc", "Phiếu xuất kho", "Xuất hàng hóa/dịch vụ theo hóa đơn bán ra"), ("02_Phieu_thu.doc", "Phiếu thu", "Thu tiền khách hàng theo hóa đơn"), ("03_Bien_ban_ban_giao_nghiem_thu.doc", "Biên bản bàn giao / nghiệm thu", "Xác nhận bàn giao hoặc nghiệm thu dịch vụ/hàng hóa"), ("04_Doi_chieu_cong_no_phai_thu.doc", "Biên bản đối chiếu công nợ phải thu", "Đối chiếu công nợ phải thu với khách hàng")]
    else:
        items = []
    for ten, title, note in items:
        docs.append((ten, _doc_bytes(title, _body_chung_tu(title, cty, ben, inv, loai, note), kieu)))
    return docs



def _checklist_57() -> list[tuple[str, str, str]]:
    groups = {
        "Folder_1_LƯƠNG_&_BHXH": ["Hợp đồng lao động", "Phụ lục hợp đồng lao động", "Quyết định tiếp nhận", "Quyết định bổ nhiệm/chức danh", "Bảng chấm công tháng", "Bảng lương tháng", "Phiếu lương/biên nhận lương", "Chứng từ chuyển khoản lương", "Đăng ký MST cá nhân", "Cam kết thu nhập mẫu 08", "Hồ sơ người phụ thuộc", "Tờ khai BHXH TK1-TS", "Danh sách D02-LT", "Chứng từ nộp BHXH", "Quyết định nghỉ việc/bàn giao"],
        "Folder_2_MUA_BÁN_&_GIA_CÔNG": ["Hóa đơn đầu vào", "Hóa đơn đầu ra", "Phiếu nhập kho", "Phiếu xuất kho", "Phiếu giao nhận", "Biên bản kiểm kê kho", "Đơn đặt hàng", "Báo giá", "Hợp đồng mua bán", "Hợp đồng nguyên tắc", "Biên bản nghiệm thu", "Biên bản thanh lý", "Đơn đặt gia công", "Hợp đồng gia công", "Biên bản nghiệm thu gia công", "Bảng kê vật tư giao gia công", "Biên bản giao nhận gia công"],
        "Folder_3_THUẾ_&_HĐĐT": ["XML hóa đơn", "PDF gốc nhà cung cấp", "Bảng kê VAT", "Tờ khai VAT", "Tờ khai TNDN", "Tờ khai TNCN", "Chứng từ nộp thuế", "Giải trình chênh lệch", "Thông báo phát hành/đăng ký hóa đơn", "Biên bản điều chỉnh/thay thế hóa đơn"],
        "Folder_4_KẾ_TOÁN_&_BÁN_HÀNG": ["Sổ nhật ký chung", "Sổ cái", "Sổ chi tiết công nợ", "Sổ quỹ tiền mặt", "Sổ tiền gửi ngân hàng", "Sao kê ngân hàng", "Phiếu thu", "Phiếu chi", "Ủy nhiệm chi", "Chứng từ POS", "Hoàn hàng/đổi trả"],
        "Folder_5_PHÁP_LÝ": ["Giấy chứng nhận đăng ký doanh nghiệp", "Điều lệ công ty", "Quyết định bổ nhiệm giám đốc", "Quyết định bổ nhiệm kế toán", "Hợp đồng thuê mặt bằng", "Cam kết/ủy quyền", "Nội quy lao động", "Quy chế lương thưởng", "Biên bản họp/quyết định chủ sở hữu", "Hồ sơ tài khoản ngân hàng", "Chữ ký số và đăng ký giao dịch điện tử"]
    }
    out = []
    for folder, names in groups.items():
        for name in names:
            out.append((folder, name, "Bắt buộc" if len(out) < 32 else "Định kỳ"))
    return out[:57]


def _mau_gia_cong(cty: dict, kieu: str) -> list[tuple[str, bytes]]:
    docs = []
    mau = {
        "01_BIEN_BAN_KIEM_KE_KHO.doc": "<h1>Biên bản kiểm kê kho</h1><p>Kiểm kê tại kho: <span class='line'></span>. Thời điểm: <span class='line'></span>.</p><table><tr><th>STT</th><th>Mã hàng</th><th>Tên hàng</th><th>ĐVT</th><th>Sổ sách</th><th>Thực tế</th><th>Chênh lệch</th><th>Ghi chú</th></tr><tr><td>1</td><td></td><td></td><td></td><td></td><td></td><td></td><td></td></tr></table><table class='sign'><tr><td>Thủ kho</td><td>Kế toán</td><td>Đại diện công ty</td></tr></table>",
        "02_PHIEU_GIAO_NHAN.doc": "<h1>Phiếu giao nhận</h1><p>Bên giao: <span class='line'></span> · Bên nhận: <span class='line'></span></p><table><tr><th>STT</th><th>Hàng hóa/dịch vụ</th><th>Số lượng</th><th>Tình trạng</th><th>Ghi chú</th></tr><tr><td>1</td><td></td><td></td><td>Đạt</td><td></td></tr></table><table class='sign'><tr><td>Bên giao</td><td>Bên nhận</td></tr></table>",
        "03_DON_DAT_GIA_CONG.doc": "<h1>Đơn đặt gia công</h1><p>Đơn vị đặt gia công: " + _html(cty['ten']) + " - MST " + _html(cty['mst']) + "</p><table><tr><th>Sản phẩm</th><th>Quy cách</th><th>Số lượng</th><th>Đơn giá gia công</th><th>Ngày giao</th></tr><tr><td></td><td></td><td></td><td></td><td></td></tr></table>",
        "04_BIEN_BAN_NGHIEM_THU_GIA_CONG.doc": "<h1>Biên bản nghiệm thu gia công</h1><p>Các bên xác nhận hàng gia công đã hoàn thành theo đơn đặt hàng/hợp đồng.</p><table><tr><th>Nội dung</th><th>Số lượng</th><th>Kết quả QC</th><th>Ghi chú</th></tr><tr><td></td><td></td><td>Đạt</td><td></td></tr></table><table class='sign'><tr><td>Bên đặt</td><td>Bên nhận gia công</td></tr></table>",
        "05_HOP_DONG_GIA_CONG.doc": "<h1>Hợp đồng gia công</h1><p>Bên đặt gia công: " + _html(cty['ten']) + " - MST " + _html(cty['mst']) + "</p><h2>Nội dung gia công</h2><p>Sản phẩm, định mức, tiêu chuẩn chất lượng, thời hạn giao nhận và phương thức thanh toán theo phụ lục/đơn đặt hàng kèm theo.</p><table class='sign'><tr><td>Bên đặt gia công</td><td>Bên nhận gia công</td></tr></table>"
    }
    for fn, body in mau.items():
        docs.append((fn, _doc_bytes(fn.rsplit('.',1)[0].replace('_',' '), body, kieu)))
    return docs

def _mau_hop_dong(cty: dict, kieu: str) -> list[tuple[str, bytes]]:
    body = f'''<p class="center"><b>CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</b><br>Độc lập - Tự do - Hạnh phúc</p><h1>Hợp đồng nguyên tắc / mua bán</h1><p class="sub">Mẫu khung tự động cho {_html(cty['ten'])}</p><h2>1. Thông tin các bên</h2><p>Bên A: <span class="line"></span> · MST: <span class="line"></span></p><p>Bên B: {_html(cty['ten'])} · MST: {_html(cty['mst'])}</p><h2>2. Nội dung hợp tác</h2><p>Các bên thỏa thuận mua bán hàng hóa/dịch vụ theo đơn đặt hàng, báo giá, biên bản giao nhận/nghiệm thu và hóa đơn điện tử phát sinh từng lần.</p><h2>3. Thanh toán và chứng từ</h2><p>Thanh toán theo từng hóa đơn hoặc bảng đối chiếu công nợ. Bộ chứng từ gồm hợp đồng/đơn đặt hàng, hóa đơn, biên bản giao nhận/nghiệm thu, đề nghị thanh toán và phiếu thu/chi hoặc UNC.</p><table class="sign"><tr><td>Đại diện Bên A</td><td>Đại diện Bên B</td></tr></table>'''
    thanh_ly = '''<h1>Biên bản nghiệm thu / thanh lý hợp đồng</h1><p>Hôm nay, ngày <span class="line"></span>, các bên xác nhận đã hoàn thành nghĩa vụ theo hợp đồng/đơn hàng số <span class="line"></span>.</p><table><tr><th>Nội dung</th><th>Kết quả nghiệm thu</th><th>Giá trị</th></tr><tr><td>Hàng hóa/dịch vụ theo hóa đơn</td><td>Đạt yêu cầu</td><td></td></tr></table><p>Các bên thống nhất công nợ còn phải thanh toán: <span class="line"></span> đồng.</p><table class="sign"><tr><td>Bên giao</td><td>Bên nhận</td></tr></table>'''
    return [("HopDongNguyenTac_MuaBan.doc", _doc_bytes("Hợp đồng nguyên tắc / mua bán", body, kieu)), ("BienBanNghiemThu_ThanhLy.doc", _doc_bytes("Biên bản nghiệm thu / thanh lý", thanh_ly, kieu))]


def _mau_lao_dong(cty: dict, kieu: str) -> list[tuple[str, bytes]]:
    hdld = f'''<h1>Hợp đồng lao động</h1><p class="sub">Mẫu khung cho {_html(cty['ten'])}</p><table><tr><td>Người sử dụng lao động</td><td>{_html(cty['ten'])} - MST {_html(cty['mst'])}</td></tr><tr><td>Người lao động</td><td><span class="line"></span></td></tr><tr><td>Chức danh / bộ phận</td><td><span class="line"></span></td></tr><tr><td>Mức lương / phụ cấp</td><td><span class="line"></span></td></tr><tr><td>Thời hạn hợp đồng</td><td><span class="line"></span></td></tr></table><p>Hai bên thỏa thuận quyền, nghĩa vụ, thời giờ làm việc, nghỉ ngơi, bảo hiểm, bảo mật và các nội dung khác theo quy định hiện hành.</p><table class="sign"><tr><td>Người lao động</td><td>Người sử dụng lao động</td></tr></table>'''
    qd = '''<h1>Quyết định tiếp nhận / điều chuyển / nghỉ việc</h1><p>Căn cứ nhu cầu sử dụng lao động và hồ sơ nhân sự, công ty quyết định:</p><table><tr><th>Họ tên</th><th>Chức danh</th><th>Ngày hiệu lực</th><th>Nội dung</th></tr><tr><td></td><td></td><td></td><td></td></tr></table><table class="sign"><tr><td>Người nhận quyết định</td><td>Người ký quyết định</td></tr></table>'''
    return [("HopDongLaoDong.doc", _doc_bytes("Hợp đồng lao động", hdld, kieu)), ("QuyetDinhNhanSu.doc", _doc_bytes("Quyết định nhân sự", qd, kieu))]


def _workbook_bytes(rows: list[dict], cty: dict, nhom: set[str]) -> bytes:
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Bang_Ke_Hoa_Don"
    ws.append(["Doanh nghiệp", cty["ten"], "MST", cty["mst"]]); ws.append([])
    ws.append(["Ngày", "Ký hiệu", "Số", "Đối tác", "MST đối tác", "Tiền hàng", "Thuế", "Tổng", "Có XML", "Có PDF gốc", "Ghi chú"])
    for inv in rows:
        b = _ben(inv); ws.append([_ngay_vn(ngay_hd(inv)), inv.get("khhdon") or inv.get("ky_hieu"), inv.get("shdon") or inv.get("so"), b["ten"], b["mst"], inv.get("tien_hang") or inv.get("tgtcthue") or 0, inv.get("thue") or inv.get("tgtthue") or 0, inv.get("tong") or inv.get("tgtttbso") or 0, "Có" if inv.get("_xml_path") else "Không", "Có" if inv.get("_goc_path") else "Không", inv.get("_goc_loi") or ""])
    for col, width in zip("ABCDEFGHIJK", [14,14,12,38,16,16,14,16,12,14,42]): ws.column_dimensions[col].width = width
    ws4 = wb.create_sheet("CHECKLIST_57")
    ws4.append(["STT", "Folder", "Chứng từ", "Mức độ", "Đã có", "Ghi chú"])
    for i, (folder, name, level) in enumerate(_checklist_57(), 1):
        ws4.append([i, folder, name, level, "☐", ""])
    ws5 = wb.create_sheet("Bang_Luong")
    ws5.append(["STT", "Họ tên", "Chức vụ", "Ngày công", "Lương cơ bản", "Phụ cấp", "Khấu trừ", "Thực nhận", "Ký nhận"])
    for i in range(1, 11): ws5.append([i, "", "", "", "", "", "", "", ""])
    ws6 = wb.create_sheet("Bang_Cham_Cong")
    ws6.append(["STT", "Họ tên"] + [str(i) for i in range(1, 32)] + ["Tổng công", "Ghi chú"])
    for i in range(1, 11): ws6.append([i, ""] + ["" for _ in range(31)] + ["", ""])
    if "thue_bhxh" in nhom:
        ws2 = wb.create_sheet("Khung_BHXH_D02LT_TK1TS"); ws2.append(["STT", "Họ tên", "Mã số BHXH", "Số CCCD", "Chức danh", "Tiền lương", "Từ tháng", "Đến tháng", "Ghi chú"])
        for i in range(1, 6): ws2.append([i, "", "", "", "", "", "", "", ""])
        ws3 = wb.create_sheet("Ho_So_Thue"); ws3.append(["Mục hồ sơ", "Đã có", "Ghi chú"])
        for name in ["Hóa đơn điện tử", "PDF gốc NCC", "Hợp đồng/đơn hàng", "Biên bản nghiệm thu/giao nhận", "Thanh toán ngân hàng/phiếu thu chi"]: ws3.append([name, "", ""])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _tao_zip_chung_tu(mst: str, p_ds: list[dict], s_ds: list[dict], nhom: set[str], kieu: str) -> bytes:
    cty = _cty(mst); buf = io.BytesIO(); all_rows = list(p_ds) + list(s_ds)
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        readme = f'''<h1>Bộ chứng từ kế toán tự động</h1><p>Doanh nghiệp: <b>{_html(cty['ten'])}</b> · MST {_html(cty['mst'])}</p><p>Gói gồm {len(p_ds)} hóa đơn mua vào, {len(s_ds)} hóa đơn bán ra. File .doc có thể mở bằng Microsoft Word/LibreOffice để rà soát, ký và in.</p><p>Lưu ý: mẫu là khung chứng từ để kế toán hoàn thiện trước khi ký/nộp; các biểu mẫu BHXH/thuế nên đối chiếu với cổng nộp hồ sơ tại thời điểm sử dụng.</p>'''
        z.writestr("00_README.html", _doc_frame("README bộ chứng từ", readme).encode("utf-8"))
        z.writestr("00_DoiSoat_HoaDon.xlsx", _workbook_bytes(all_rows, cty, nhom))
        z.writestr("00_manifest.json", json.dumps({"mst": mst, "ten_doanh_nghiep": cty["ten"], "so_hoa_don_mua_vao": len(p_ds), "so_hoa_don_ban_ra": len(s_ds), "nhom_mau": sorted(nhom), "tao_luc": dt.datetime.now().isoformat(timespec="seconds")}, ensure_ascii=False, indent=2))
        for l, rows, root in (("purchase", p_ds, "01_MuaVao"), ("sold", s_ds, "02_BanRa")):
            for inv in rows:
                ten = f"{_ngay_vn(ngay_hd(inv)).replace('/', '-')}_{_slug(_hoa_don_id(inv))}_{_slug(_ben(inv)['ten'], 40)}"
                folder = f"{root}/{ten}"
                for fn, data in _invoice_docs(cty, inv, l, nhom, kieu): z.writestr(f"{folder}/{fn}", data)
                z.writestr(f"{folder}/ThongTin_HoaDon.json", json.dumps(tom_tat(inv), ensure_ascii=False, indent=2))
        if "hop_dong" in nhom:
            for fn, data in _mau_gia_cong(cty, kieu): z.writestr("02_MuaBan_GiaCong/" + fn, data)
            for fn, data in _mau_hop_dong(cty, kieu): z.writestr("02_MuaBan_GiaCong/" + fn, data)
        if "lao_dong" in nhom:
            for fn, data in _mau_lao_dong(cty, kieu): z.writestr("01_Luong_BHXH/" + fn, data)
        if "thue_bhxh" in nhom:
            body = '''<h1>Checklist hồ sơ thuế và BHXH</h1><table><tr><th>Nhóm hồ sơ</th><th>Tệp cần chuẩn bị</th><th>Trạng thái</th></tr><tr><td>Thuế GTGT</td><td>Bảng kê hóa đơn, hóa đơn XML/PDF, giải trình chênh lệch</td><td>Rà soát</td></tr><tr><td>BHXH</td><td>D02-LT/TK1-TS, HĐLĐ, quyết định lương, CCCD</td><td>Điền bổ sung nhân sự</td></tr></table>'''
            z.writestr("03_Thue_HDDT/Checklist_HoSo.doc", _doc_bytes("Checklist hồ sơ thuế và BHXH", body, kieu))
        for folder, desc in [("01_Luong_BHXH/", "Lương, chấm công, BHXH"), ("02_MuaBan_GiaCong/", "Mua bán, kho và gia công"), ("03_Thue_HDDT/", "Thuế và hóa đơn điện tử"), ("04_KeToan_BanHang/", "Sổ kế toán, ngân hàng, POS, hoàn hàng"), ("05_PhapLy/", "Hồ sơ pháp lý doanh nghiệp")]:
            z.writestr(folder + "README.txt", desc)
    return buf.getvalue()

# ------------------------------------------------------------------ cửa vào

def ly_do_chan(app_id: str) -> Optional[str]:
    ten = "Tải Hoá Đơn TCT" if app_id == APP_TAI else "Quét hoá đơn Gmail" if app_id == APP_GMAIL else "Bộ chứng từ kế toán"
    if not giay_phep.da_mo(app_id):
        return f"{ten} chưa được mở khoá. Vào Kết nối & Khách > Kết nối ngoài để nhập mã kích hoạt."
    if not knn.da_dong_y("chung"):
        return "Cần đồng ý Điều khoản sử dụng MinMin trước (trang Kết nối ngoài)."
    return None


def _chan(app_id: str) -> Optional[JSONResponse]:
    ly = ly_do_chan(app_id)
    return JSONResponse({"ok": False, "error": ly}, status_code=403) if ly else None


def _chan_chung_tu() -> Optional[JSONResponse]:
    if ly_do_chan(APP_TAI) is None or ly_do_chan(APP_CT) is None:
        return None
    return JSONResponse({"ok": False, "error": ly_do_chan(APP_CT)}, status_code=403)


def _loi(e: Exception, ma: int = 400) -> JSONResponse:
    return JSONResponse({"ok": False, "error": str(e)}, status_code=ma)


async def _than(request: Request) -> dict:
    try:
        d = await request.json()
        return d if isinstance(d, dict) else {}
    except Exception:
        return {}


def _tai_ve(data: bytes, ten: str, kieu: str) -> Response:
    from urllib.parse import quote
    return Response(data, media_type=kieu, headers={
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(ten)}",
        "Cache-Control": "no-store"})


def _trang_khoa(ly_do: str) -> HTMLResponse:
    from html import escape
    return HTMLResponse(
        '<!doctype html><html lang="vi"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1"><title>Hoá đơn</title>'
        '<style>body{font-family:system-ui,sans-serif;background:#0f1420;color:#e8ecf4;display:grid;'
        'place-items:center;min-height:100vh;margin:0;padding:16px}div{max-width:460px;background:#182033;'
        'border-radius:14px;padding:24px;line-height:1.55}a{color:#7cc4ff}</style></head><body><div>'
        f'<h2>Hoá đơn</h2><p>{escape(ly_do)}</p><p><a href="/">Về MinMin</a></p></div></body></html>',
        status_code=403)


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def register(app: FastAPI) -> None:
    A = GOC + "/api"

    # ---------------- giao diện
    @app.get(GOC)
    def hd_goc():
        return RedirectResponse(GOC + "/", status_code=307)

    @app.get(GOC + "/")
    def hd_trang():
        if ly_do_chan(APP_TAI) and ly_do_chan(APP_GMAIL) and ly_do_chan(APP_CT):
            return _trang_khoa(ly_do_chan(APP_TAI))
        return FileResponse(str(WEB_DIR / "index.html"), headers={"Cache-Control": "no-cache"})

    @app.get(A + "/trang-thai")
    def hd_trang_thai():
        try:
            co_cffi = bool(thd("gdt_client")._HAS_CFFI)
        except Exception:
            co_cffi = False
        return {"ok": True, "tai_hoa_don": ly_do_chan(APP_TAI) is None,
                "gmail": ly_do_chan(APP_GMAIL) is None,
                "bo_chung_tu": (ly_do_chan(APP_TAI) is None or ly_do_chan(APP_CT) is None),
                "ly_do_tai": ly_do_chan(APP_TAI), "ly_do_gmail": ly_do_chan(APP_GMAIL),
                "ly_do_bo_chung_tu": ly_do_chan(APP_CT),
                "gia_lap_chrome": co_cffi}

    # ---------------- việc nền (dùng chung)
    @app.get(A + "/viec/{vid}")
    def hd_viec(vid: str, tu: int = 0):
        v = VIEC.get(vid)
        if not v:
            return _loi(ValueError("Không có việc này (có thể MinMin vừa khởi động lại)."), 404)
        return {"ok": True, **v.ra_dict(max(0, tu))}

    @app.post(A + "/viec/{vid}/dung")
    def hd_viec_dung(vid: str):
        v = VIEC.get(vid)
        if v:
            v.huy.set()
        return {"ok": bool(v)}

    # ---------------- doanh nghiệp + đăng nhập
    @app.get(A + "/doanh-nghiep")
    def hd_dn():
        if (c := _chan(APP_TAI)):
            return c
        out = []
        for mst, d in sorted(doc_cau_hinh()["doanh_nghiep"].items()):
            ok, exp = da_dang_nhap(mst)
            out.append({"mst": mst, "ten": d.get("ten", ""), "co_mk": bool(d.get("mk")),
                        "dang_nhap": ok, "het_han": exp,
                        "so_hd": {l: len(doc_kho(mst, l)) for l in LOAI}})
        return {"ok": True, "ds": out}

    @app.post(A + "/doanh-nghiep")
    async def hd_dn_luu(request: Request):
        if (c := _chan(APP_TAI)):
            return c
        d = await _than(request)
        try:
            mst = chuan_mst(d.get("mst"))
        except ValueError as e:
            return _loi(e)
        ch = doc_cau_hinh()
        dn = ch["doanh_nghiep"].setdefault(mst, {})
        if d.get("ten") is not None:
            dn["ten"] = str(d.get("ten") or "")[:200]
        if d.get("mat_khau"):
            dn["mk"] = ma_hoa(str(d["mat_khau"]))
        ghi_cau_hinh(ch)
        return {"ok": True, "mst": mst}

    @app.delete(A + "/doanh-nghiep/{mst}")
    def hd_dn_xoa(mst: str):
        if (c := _chan(APP_TAI)):
            return c
        ch = doc_cau_hinh()
        ch["doanh_nghiep"].pop(mst, None)
        ghi_cau_hinh(ch)
        _CLIENT.pop(mst, None)
        try:
            (goc_du_lieu() / "phien" / f"{chuan_mst(mst)}.json").unlink()
        except (OSError, ValueError):
            pass
        return {"ok": True}

    @app.get(A + "/captcha")
    def hd_captcha(mst: str):
        if (c := _chan(APP_TAI)):
            return c
        try:
            cap = client_cho(mst).get_captcha()
        except Exception as e:
            return _loi(e, 502)
        svg = thd("captcha_image").normalize_svg(str(cap.get("content") or ""))
        return {"ok": True, "key": cap.get("key", ""), "svg": svg}

    @app.post(A + "/dang-nhap")
    async def hd_dang_nhap(request: Request):
        if (c := _chan(APP_TAI)):
            return c
        d = await _than(request)
        try:
            mst = chuan_mst(d.get("mst"))
            mk = str(d.get("mat_khau") or "") or giai_ma(doc_cau_hinh()["doanh_nghiep"].get(mst, {}).get("mk", ""))
            if not mk:
                raise ValueError("Chưa có mật khẩu cổng thuế của MST này.")
            cl = client_cho(mst)
            cl.login(mst, mk, str(d.get("captcha") or "").strip(), str(d.get("ckey") or ""))
        except Exception as e:
            return _loi(e)
        luu_phien(mst, cl)
        if d.get("nho_mk") and d.get("mat_khau"):
            ch = doc_cau_hinh()
            ch["doanh_nghiep"].setdefault(mst, {})["mk"] = ma_hoa(str(d["mat_khau"]))
            ghi_cau_hinh(ch)
        else:
            ch = doc_cau_hinh()
            ch["doanh_nghiep"].setdefault(mst, {})
            ghi_cau_hinh(ch)
        info = cl.token_info()
        return {"ok": True, "het_han": cl.token_exp(), "ten": info.get("name") or ""}

    @app.post(A + "/dang-xuat/{mst}")
    def hd_dang_xuat(mst: str):
        if (c := _chan(APP_TAI)):
            return c
        _CLIENT.pop(mst, None)
        try:
            (goc_du_lieu() / "phien" / f"{chuan_mst(mst)}.json").unlink()
        except (OSError, ValueError):
            pass
        return {"ok": True}

    # ---------------- tra cứu + tải
    @app.post(A + "/tra-cuu")
    async def hd_tra_cuu(request: Request):
        if (c := _chan(APP_TAI)):
            return c
        d = await _than(request)
        try:
            mst = chuan_mst(d.get("mst"))
            tu, den = _ngay(d.get("tu")), _ngay(d.get("den"))
            if den < tu:
                raise ValueError("Ngày kết thúc trước ngày bắt đầu.")
            if (den - tu).days > 366:
                raise ValueError("Mỗi lần tra tối đa 12 tháng.")
            loai = d.get("loai") or "ca_hai"
            loai_ds = list(LOAI) if loai == "ca_hai" else [loai]
            if not set(loai_ds) <= set(LOAI):
                raise ValueError("Loại hoá đơn không hợp lệ.")
            he = d.get("he") or "query"
            he_ds = ["query", "sco-query"] if he == "ca_hai" else [he]
            if not set(he_ds) <= {"query", "sco-query"}:
                raise ValueError("Hệ thống hoá đơn không hợp lệ.")
            _can_dang_nhap(mst)
            v = chay_viec("tra-cuu", f"Tra cứu {mst} {tu:%d/%m/%Y} - {den:%d/%m/%Y}",
                          lambda v: tra_cuu(v, mst, tu, den, loai_ds, he_ds))
        except Exception as e:
            return _loi(e)
        return {"ok": True, "viec": v.id}

    @app.get(A + "/hoa-don")
    def hd_ds(mst: str, loai: str = "purchase", tu: str = "", den: str = ""):
        if (c := _chan(APP_TAI)):
            return c
        if loai not in LOAI:
            return _loi(ValueError("Loại hoá đơn không hợp lệ."))
        try:
            ds = loc_kho(mst, loai, tu, den)
        except ValueError as e:
            return _loi(e)
        rows = [tom_tat(x) for x in ds]
        tong = {"so_hd": len(rows), "tien_hang": sum(float(r["tien_hang"] or 0) for r in rows),
                "thue": sum(float(r["thue"] or 0) for r in rows), "tong": sum(float(r["tong"] or 0) for r in rows),
                "co_xml": sum(1 for r in rows if r["co_xml"]), "co_goc": sum(1 for r in rows if r["co_goc"])}
        return {"ok": True, "ds": rows, "tong": tong}

    @app.post(A + "/tai")
    async def hd_tai(request: Request):
        if (c := _chan(APP_TAI)):
            return c
        d = await _than(request)
        try:
            mst = chuan_mst(d.get("mst"))
            loai = d.get("loai")
            if loai not in LOAI:
                raise ValueError("Loại hoá đơn không hợp lệ.")
            keys = d.get("keys")
            if keys is None:
                keys = [khoa_hd(x) for x in loc_kho(mst, loai, d.get("tu"), d.get("den"))]
            if not isinstance(keys, list) or not keys or len(keys) > 5000:
                raise ValueError("Chọn từ 1 đến 5000 hoá đơn mỗi lượt tải.")
            keys = list(dict.fromkeys(str(k) for k in keys))
            kieu = [k for k in (d.get("kieu") or []) if k in ("xml", "pdf", "goc")]
            if d.get("ban_goc") and "goc" not in kieu:      # tương thích giao diện cũ
                kieu.append("goc")
            if not kieu:
                kieu = ["xml"]
            if "xml" not in kieu:
                kieu.insert(0, "xml")                         # PDF nào cũng cần XML trước
            _can_dang_nhap(mst)
            v = chay_viec("tai", f"Tải {len(keys)} hoá đơn {mst}",
                          lambda v: tai_ds(v, mst, loai, keys, kieu))
        except Exception as e:
            return _loi(e)
        return {"ok": True, "viec": v.id, "so": len(keys)}

    def _tim(mst: str, loai: str, key: str) -> Optional[dict]:
        if loai not in LOAI:
            return None
        return doc_kho(mst, loai).get(key)

    @app.get(A + "/xem/{mst}/{loai}/{key}")
    def hd_xem(mst: str, loai: str, key: str):
        if (c := _chan(APP_TAI)):
            return c
        try:
            inv = _tim(mst, loai, key)
        except ValueError:
            inv = None
        if not inv:
            return HTMLResponse("Không có hoá đơn này.", status_code=404)
        if inv.get("_detail"):
            return HTMLResponse(thd("invoice_xml").render_html(inv["_detail"]),
                                headers={"Content-Security-Policy": "default-src 'none'; style-src 'unsafe-inline'"})
        return HTMLResponse("Hoá đơn này chưa tải XML. Bấm Tải trước rồi xem lại.", status_code=404)

    @app.get(A + "/tep/{mst}/{loai}/{key}/{kieu}")
    def hd_tep(mst: str, loai: str, key: str, kieu: str):
        if (c := _chan(APP_TAI)):
            return c
        try:
            inv = _tim(mst, loai, key)
        except ValueError:
            inv = None
        truong = {"xml": "_xml_path", "goc": "_goc_path", "html": "_html_path", "zip": "_zip_path",
                  "pdf": "_pdf_path"}.get(kieu)
        p = tuyet_doi(inv.get(truong)) if (inv and truong) else None
        if not p:
            return _loi(ValueError("Chưa có file này."), 404)
        kieu_mime = {"xml": "application/xml", "goc": "application/pdf", "html": "text/html",
                     "zip": "application/zip", "pdf": "application/pdf"}[kieu]
        return _tai_ve(p.read_bytes(), p.name, kieu_mime)

    # ---------------- cài đặt tải
    def _mo_ta_cai_dat() -> dict:
        cd = cai_dat_tai()
        mau = {"tdlap": "2026-09-15T00:00:00", "shdon": 123, "khmshdon": 1, "khhdon": "C26TAA",
               "nbmst": "0312345678", "nbten": "CÔNG TY TNHH THỜI TRANG ÁNH DƯƠNG", "_loai": "purchase"}
        return {"ok": True, "cai_dat": cd,
                "token_ten_file": [{"id": a, "ten": b} for a, b in TOKEN_TEN_FILE],
                "token_ten_excel": [{"id": a, "ten": b} for a, b in TOKEN_TEN_EXCEL],
                "xep_thu_muc": [{"id": a, "ten": b} for a, b in XEP_THU_MUC.items()],
                "vi_du_ten_file": ten_file(mau, cd), "vi_du_ten_excel": ten_excel("0312345678", "purchase",
                                                                                   "2026-09-01", "2026-09-30", cd),
                "vi_du_thu_muc": "/".join(_thu_muc_tai("0312345678", mau, cd).parts[-3:]),
                "co_trinh_duyet": bool(trinh_duyet_in_pdf())}

    @app.get(A + "/cai-dat")
    def hd_cai_dat():
        if (c := _chan(APP_TAI)):
            return c
        return _mo_ta_cai_dat()

    @app.get(A + "/captcha-api")
    def cap_doc():
        if ly_do_chan(APP_TAI) and ly_do_chan(APP_GMAIL):
            return _chan(APP_TAI)
        return {"ok": True, **captcha_cong_khai()}

    @app.post(A + "/captcha-api")
    async def cap_luu(request: Request):
        if ly_do_chan(APP_TAI) and ly_do_chan(APP_GMAIL):
            return _chan(APP_TAI)
        d = await _than(request)
        p = d.get("provider")
        if p not in ("2captcha", "omocaptcha"):
            return _loi(ValueError("Chọn 2Captcha hoặc OmoCaptcha."))
        key = str(d.get("api_key") or "").strip()
        if len(key) > 512:
            return _loi(ValueError("API key quá dài."))
        with _KHOA_GHI:
            ch = doc_cau_hinh()
            c = ch.setdefault("captcha", {})
            if d.get("xoa_key"):
                c.pop(p, None)
            elif key:
                c[p] = ma_hoa(key)
            c.update(provider=p, enabled=bool(d.get("enabled")))
            if c["enabled"] and not giai_ma(c.get(p, "")):
                return _loi(ValueError("Nhập API key trước khi bật."))
            ghi_cau_hinh(ch)
        return {"ok": True, **captcha_cong_khai()}

    @app.post(A + "/tra-cuu-hang")
    async def hd_hang(request: Request):
        if (c := _chan(APP_TAI)):
            return c
        d = await _than(request)
        try:
            mst, loai = chuan_mst(d.get("mst")), d.get("loai")
            if loai not in LOAI:
                raise ValueError("Loại hoá đơn không hợp lệ.")
            with _KHOA_GHI:
                kho = doc_kho(mst, loai)
                inv = kho.get(str(d.get("key")))
                if inv is None:
                    raise ValueError("Không tìm thấy hoá đơn.")
                if "code" in d:
                    code = str(d["code"]).strip()
                    if not re.fullmatch(r"[A-Za-z0-9._-]{0,200}", code):
                        raise ValueError("Mã tra cứu chỉ gồm chữ, số, dấu chấm, gạch ngang/gạch dưới.")
                    inv["_ma_tra_cuu"] = code
                    inv["_bkav_mtc"] = code
                    inv["_code"] = code
                    ghi_kho(mst, loai, kho)
                return {"ok": True, **ma_tra_cuu_hd(inv)}
        except Exception as e:
            return _loi(e)

    @app.post(A + "/captcha-api/so-du")
    async def cap_so_du(request: Request):
        if ly_do_chan(APP_TAI) and ly_do_chan(APP_GMAIL):
            return _chan(APP_TAI)
        d = await _than(request)
        p = d.get("provider")
        if p not in ("2captcha", "omocaptcha"):
            return _loi(ValueError("Dịch vụ không hợp lệ."))
        key = giai_ma(doc_cau_hinh().get("captcha", {}).get(p, ""))
        if not key:
            return _loi(ValueError("Lưu API key dịch vụ này trước."))
        try:
            from starlette.concurrency import run_in_threadpool
            balance = await run_in_threadpool(thd("captcha_solver").get_balance, p, key)
            return {"ok": True, "balance": balance}
        except Exception:
            return _loi(ValueError("Chưa kiểm tra được số dư. Kiểm tra API key và kết nối tới dịch vụ."))

    @app.post(A + "/cai-dat")
    async def hd_cai_dat_luu(request: Request):
        if (c := _chan(APP_TAI)):
            return c
        d = await _than(request)
        ch = doc_cau_hinh()
        cu = ch.get("tai") or {}
        for k in ("ten_file", "ten_excel", "xep_thu_muc", "giu_zip", "luong"):
            if k in d:
                cu[k] = d[k]
        ch["tai"] = cu
        ghi_cau_hinh(ch)
        return _mo_ta_cai_dat()

    def _chon(mst: str, d: dict) -> tuple[list, list, str, str]:
        """Hoá đơn được xuất: các dòng đã tích (keys + loai) hoặc toàn bộ trong khoảng ngày."""
        tu, den = (d.get("tu") or "")[:10], (d.get("den") or "")[:10]
        keys = d.get("keys") if isinstance(d.get("keys"), list) else []
        if "keys" in d and d.get("loai") in LOAI:
            kho = doc_kho(mst, d["loai"])
            ds = [kho[k] for k in keys if k in kho]
            return (ds, []) if d["loai"] == "purchase" else ([], ds), tu, den
        return (loc_kho(mst, "purchase", tu, den), loc_kho(mst, "sold", tu, den)), tu, den

    @app.post(A + "/xuat")
    async def hd_xuat(request: Request):
        if (c := _chan(APP_TAI)):
            return c
        d = await _than(request)
        try:
            mst = chuan_mst(d.get("mst"))
            (p_ds, s_ds), tu, den = _chon(mst, d)
            if not (p_ds or s_ds):
                raise ValueError("Chưa có hoá đơn nào để xuất (chọn dòng, hoặc tra cứu khoảng ngày trước).")
            ex = thd("exporter")
            kieu = d.get("kieu") or "excel"
            cd = cai_dat_tai()
            loai_ten = "purchase" if p_ds and not s_ds else "sold" if s_ds and not p_ds else ""
            nhan = ten_excel(mst, loai_ten, tu, den, cd)
            tat_ca = p_ds + s_ds
            if kieu in ("excel", "bang_ke", "chi_tiet"):
                # exporter chỉ kiểm có/không của đường XML: đưa đường tuyệt đối vào BẢN SAO
                ban_sao = [dict(x, _xml_path=str(tuyet_doi(x.get("_xml_path")) or "") or None) for x in tat_ca]
                p2, s2 = ban_sao[:len(p_ds)], ban_sao[len(p_ds):]
                import tempfile
                with tempfile.TemporaryDirectory() as td:
                    f = os.path.join(td, "x.xlsx")
                    if kieu == "excel":
                        ex.export_excel(f, p2, s2)
                        ten = nhan
                    elif kieu == "bang_ke":
                        ky = f"{tu[8:10]}/{tu[5:7]}/{tu[:4]} - {den[8:10]}/{den[5:7]}/{den[:4]}" if tu and den else ""
                        ex.export_tax_declaration(f, p2, s2, mst_cty=mst, ky=ky)
                        ten = "BangKe_GTGT_" + nhan
                    else:
                        ex.export_detail(f, p2, s2)
                        ten = "ChiTietMatHang_" + nhan
                    return _tai_ve(Path(f).read_bytes(), ten + ".xlsx", XLSX)
            if kieu in ("zip", "html_zip"):
                loai_tep = ["html"] if kieu == "html_zip" else \
                    [t for t in (d.get("loai_tep") or ["xml", "pdf", "goc"]) if t in LOAI_TEP]
                truong = {"xml": "_xml_path", "pdf": "_pdf_path", "goc": "_goc_path", "html": "_html_path",
                          "zip": "_zip_path"}
                buf = io.BytesIO()
                n = 0
                doi_soat = []
                with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
                    for inv in tat_ca:
                        for t in loai_tep:
                            f = tuyet_doi(inv.get(truong[t]))
                            if f:
                                # giữ nguyên cây thư mục như trên máy chủ: MuaVao/2026-09/...
                                try:
                                    ten_trong = str(f.relative_to(_thu_muc_mst(mst)))
                                except ValueError:
                                    ten_trong = f.name
                                z.write(f, ten_trong)
                                n += 1
                            doi_soat.append({"hoa_don": khoa_hd(inv), "loai_file": t,
                                             "trang_thai": "da_co" if f else "con_thieu",
                                             "file": ten_trong if f else "",
                                             "ly_do": "" if f else inv.get(f"_{t}_loi", "Chưa tải được file")})
                    z.writestr("doi-soat-file.json", json.dumps(doi_soat, ensure_ascii=False, indent=2))
                if not n:
                    raise ValueError("Chưa có file loại này cho các hoá đơn đã chọn. Bấm Tải trước.")
                ten = ("HTML_TrongDong_" if kieu == "html_zip" else "HoaDon_") + nhan
                return _tai_ve(buf.getvalue(), ten + ".zip", "application/zip")
            if kieu == "pdf_gop":
                che_do = d.get("che_do") or "goc_uu_tien"
                ds_pdf = []
                for inv in tat_ca:
                    g, t = tuyet_doi(inv.get("_goc_path")), tuyet_doi(inv.get("_pdf_path"))
                    f = (g or t) if che_do == "goc_uu_tien" else g if che_do == "goc" else t
                    if f and f.read_bytes()[:5].startswith(b"%PDF"):
                        ds_pdf.append(str(f))
                if not ds_pdf:
                    raise ValueError("Chưa có PDF nào để gộp. Tải PDF Trống đồng hoặc PDF gốc trước.")
                import tempfile
                with tempfile.TemporaryDirectory() as td:
                    out = os.path.join(td, "gop.pdf")
                    so = ex.merge_pdfs(ds_pdf, out)
                    return _tai_ve(Path(out).read_bytes(), f"IN_GOP_{so}HD_{nhan}.pdf", "application/pdf")
            raise ValueError("Kiểu xuất không hợp lệ.")
        except Exception as e:
            return _loi(e)


    @app.post(A + "/bo-chung-tu")
    async def hd_bo_chung_tu(request: Request):
        if (c := _chan_chung_tu()):
            return c
        d = await _than(request)
        try:
            mst = chuan_mst(d.get("mst"))
            (p_ds, s_ds), tu, den = _chon(mst, d)
            if not (p_ds or s_ds):
                raise ValueError("Chưa có hoá đơn nào để tạo bộ chứng từ.")
            nhom = set(x for x in (d.get("nhom") or []) if x in {"mua_vao", "ban_ra", "hop_dong", "lao_dong", "thue_bhxh"})
            if not nhom:
                raise ValueError("Chọn ít nhất một bộ hồ sơ cần tạo.")
            kieu = d.get("kieu_mau") if d.get("kieu_mau") in {"ke_toan", "trang_trong", "toi_gian"} else "ke_toan"
            data = _tao_zip_chung_tu(mst, p_ds, s_ds, nhom, kieu)
            nhan = ten_excel(mst, "", tu, den, cai_dat_tai()).replace(".xlsx", "")
            return _tai_ve(data, "BoChungTu_" + nhan + ".zip", "application/zip")
        except Exception as e:
            return _loi(e)

    @app.get(A + "/canh-bao")
    def hd_canh_bao(mst: str, tu: str = "", den: str = ""):
        if (c := _chan(APP_TAI)):
            return c
        try:
            ds = thd("analysis").check_invoices(loc_kho(mst, "purchase", tu, den),
                                                loc_kho(mst, "sold", tu, den), chuan_mst(mst))
        except Exception as e:
            return _loi(e)
        out = [{k: w.get(k) for k in ("muc", "loai", "ma", "so", "khhdon", "ngay", "doi_tac", "noi_dung")}
               for w in ds[:500]]
        return {"ok": True, "ds": out}

    # ---------------- Gmail
    def _goc_web(request: Request) -> str:
        """Địa chỉ gốc mà TRÌNH DUYỆT thấy (sau Caddy/nginx), để Google trả người dùng về đúng chỗ."""
        proto = (request.headers.get("x-forwarded-proto") or request.url.scheme).split(",")[0].strip()
        host = (request.headers.get("x-forwarded-host") or request.headers.get("host") or request.url.netloc).split(",")[0].strip()
        return f"{proto}://{host}"

    def _ve_google(request: Request) -> str:
        return _goc_web(request) + A + "/gmail/google/ve"

    @app.get(A + "/gmail/cau-hinh")
    def gm_ch(request: Request):
        if (c := _chan(APP_GMAIL)):
            return c
        ch = doc_cau_hinh()
        ds = [{"email": e, "kieu": t.get("kieu", "mk"),
               "co_moc": bool(ch["gmail_moc"].get(("google:" if t.get("kieu") == "google" else "") + e))}
              for e, t in sorted(ds_gmail(ch).items())]
        cid, sec = google_client()
        return {"ok": True, "ds": ds, "google": {"co_client": bool(cid and sec), "client_id": cid,
                                                 "duong_ve": _ve_google(request)}}

    @app.post(A + "/gmail/cau-hinh")
    async def gm_ch_luu(request: Request):
        """Thêm / sửa tài khoản dùng MẬT KHẨU ỨNG DỤNG."""
        if (c := _chan(APP_GMAIL)):
            return c
        d = await _than(request)
        mail = str(d.get("email") or "").strip().lower()
        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", mail):
            return _loi(ValueError("Địa chỉ Gmail không hợp lệ."))
        mk = str(d.get("mat_khau") or "").replace(" ", "")
        cu = ds_gmail().get(mail) or {}
        if not mk and cu.get("kieu") != "mk":
            return _loi(ValueError("Nhập mật khẩu ứng dụng 16 ký tự."))
        if mk and not re.fullmatch(r"[a-zA-Z]{16}", mk):
            return _loi(ValueError("Mật khẩu ứng dụng là đúng 16 chữ cái Google cấp (vd: abcd efgh ijkl mnop), "
                                   "không phải mật khẩu Gmail. Tài khoản đăng nhập bằng g.co/sc hay mã PIN thì bấm "
                                   "Đăng nhập bằng Google."))
        luu_gmail(mail, {"kieu": "mk", "mk": ma_hoa(mk) if mk else cu.get("mk", "")})
        return {"ok": True}

    @app.delete(A + "/gmail/tai-khoan/{email}")
    def gm_xoa(email: str):
        if (c := _chan(APP_GMAIL)):
            return c
        tk = ds_gmail().get(email)
        if tk and tk.get("kieu") == "google":
            # Thu hồi quyền ở phía Google luôn, không để token mồ côi
            try:
                import requests
                requests.post("https://oauth2.googleapis.com/revoke",
                               data={"token": giai_ma(tk.get("refresh", ""))}, timeout=15)
            except Exception:
                pass
        luu_gmail(email, None)
        return {"ok": True}

    @app.post(A + "/gmail/google/client")
    async def gm_google_client(request: Request):
        if (c := _chan(APP_GMAIL)):
            return c
        d = await _than(request)
        cid = str(d.get("client_id") or "").strip()
        sec = str(d.get("client_secret") or "").strip()
        if not re.fullmatch(r"[0-9A-Za-z\-_.]+\.apps\.googleusercontent\.com", cid):
            return _loi(ValueError("Client ID phải có dạng ....apps.googleusercontent.com"))
        ch = doc_cau_hinh()
        cu = ch.get("google_oauth") or {}
        if not sec and not cu.get("client_secret"):
            return _loi(ValueError("Nhập Client secret."))
        ch["google_oauth"] = {"client_id": cid, "client_secret": ma_hoa(sec) if sec else cu["client_secret"]}
        ghi_cau_hinh(ch)
        return {"ok": True}

    @app.get(A + "/gmail/google/bat-dau")
    def gm_google_bat_dau(request: Request):
        if (ly := ly_do_chan(APP_GMAIL)):
            return _trang_khoa(ly)
        cid, sec = google_client()
        if not (cid and sec):
            return _trang_khoa("Chưa cài Google OAuth Client. Quay lại trang Quét hoá đơn Gmail, mở mục Đăng nhập bằng Google.")
        import secrets as _sec
        from urllib.parse import urlencode
        st = _sec.token_urlsafe(24)
        now = time.time()
        for k in [k for k, (t, _) in _OAUTH_STATE.items() if now - t > 900]:
            _OAUTH_STATE.pop(k, None)
        _OAUTH_STATE[st] = (now, _ve_google(request))
        q = urlencode({"client_id": cid, "redirect_uri": _ve_google(request), "response_type": "code",
                       "scope": GMAIL_SCOPE + " email", "access_type": "offline", "prompt": "consent select_account",
                       "include_granted_scopes": "true", "state": st})
        return RedirectResponse(GOOGLE_AUTH + "?" + q, status_code=302)

    @app.get(A + "/gmail/google/ve")
    def gm_google_ve(state: str = "", code: str = "", error: str = ""):
        if (ly := ly_do_chan(APP_GMAIL)):
            return _trang_khoa(ly)
        luu = _OAUTH_STATE.pop(state, None)
        if not luu or time.time() - luu[0] > 900:
            return _trang_khoa("Phiên đăng nhập Google đã hết hoặc không hợp lệ. Bấm Đăng nhập bằng Google lại.")
        if error or not code:
            return _trang_khoa("Bố đã huỷ hoặc Google từ chối cấp quyền" + (f" ({error})." if error else "."))
        import requests
        cid, sec = google_client()
        try:
            r = requests.post(GOOGLE_TOKEN, data={"code": code, "client_id": cid, "client_secret": sec,
                                                  "redirect_uri": luu[1], "grant_type": "authorization_code"}, timeout=30)
            d = r.json()
            if r.status_code != 200 or not d.get("access_token"):
                raise RuntimeError(d.get("error_description") or d.get("error") or f"HTTP {r.status_code}")
            if GMAIL_SCOPE not in str(d.get("scope", "")):
                raise RuntimeError("Chưa tích ô cho phép đọc Gmail ở trang Google. Đăng nhập lại và tích ô đó.")
            if not d.get("refresh_token"):
                raise RuntimeError("Google không trả quyền dài hạn. Vào myaccount.google.com/permissions gỡ MinMin rồi đăng nhập lại.")
            hs = requests.get(GMAIL_API + "/profile", headers={"Authorization": "Bearer " + d["access_token"]}, timeout=30).json()
            mail = str(hs.get("emailAddress") or "").lower()
            if not mail:
                raise RuntimeError("Không đọc được địa chỉ Gmail.")
        except Exception as e:
            return _trang_khoa(f"Đăng nhập Google chưa xong: {e}")
        luu_gmail(mail, {"kieu": "google", "refresh": ma_hoa(d["refresh_token"])})
        return RedirectResponse(GOC + "/?tab=gmail&them=" + mail, status_code=302)

    @app.post(A + "/gmail/quet")
    async def gm_quet(request: Request):
        if (c := _chan(APP_GMAIL)):
            return c
        d = await _than(request)
        try:
            tu, den = _ngay(d.get("tu")), _ngay(d.get("den"))
            if den < tu:
                raise ValueError("Ngày kết thúc trước ngày bắt đầu.")
            chi_moi = bool(d.get("chi_moi"))
            email = str(d.get("email") or "")
            ds = ds_gmail()
            if not ds:
                raise ValueError("Chưa thêm tài khoản Gmail nào.")
            if email and email not in ds:
                raise ValueError("Không có tài khoản Gmail này.")
            ten = email or next(iter(ds))
            v = chay_viec("gmail", f"Quét {ten} {tu:%d/%m/%Y} - {den:%d/%m/%Y}",
                          lambda v: quet_gmail(v, tu, den, chi_moi, ten))
        except Exception as e:
            return _loi(e)
        return {"ok": True, "viec": v.id}

    @app.get(A + "/gmail/ket-qua")
    def gm_kq():
        if (c := _chan(APP_GMAIL)):
            return c
        ds = sorted(_doc_json(_gmail_kho_path(), {}).values(), key=lambda r: r.get("luc", 0), reverse=True)
        kho_tct = {}
        for r in ds:
            thong_tin_tra_cuu_gmail(r)
            r["co_xml"] = bool(tuyet_doi(r.get("xml")))
            r["co_pdf"] = bool(tuyet_doi(r.get("pdf")))
            r["co_goc"] = r["co_pdf"]
            r["co_trong_dong"] = bool(gmail_trong_dong(r, kho_tct))
        return {"ok": True, "ds": ds}

    @app.get(A + "/gmail/tep/{rid}/{kieu}")
    def gm_tep(rid: str, kieu: str):
        if (c := _chan(APP_GMAIL)):
            return c
        r = _doc_json(_gmail_kho_path(), {}).get(rid) or {}
        p = gmail_trong_dong(r) if kieu == "trong_dong" else tuyet_doi(r.get(kieu)) if kieu in ("xml", "pdf") else None
        if not p:
            return _loi(ValueError("Chưa có file này."), 404)
        return _tai_ve(p.read_bytes(), p.name, "application/pdf" if kieu in ("pdf", "trong_dong") else "application/xml")

    @app.get(A + "/gmail/xem/{rid}")
    def gm_xem(rid: str):
        if (c := _chan(APP_GMAIL)):
            return c
        r = _doc_json(_gmail_kho_path(), {}).get(rid) or {}
        p = tuyet_doi(r.get("xml"))
        if not p:
            return HTMLResponse("Email này không đính kèm XML.", status_code=404)
        ixml = thd("invoice_xml")
        return HTMLResponse(ixml.render_html(ixml.parse_invoice(p.read_bytes())),
                            headers={"Content-Security-Policy": "default-src 'none'; style-src 'unsafe-inline'"})

    @app.post(A + "/gmail/ghep-lai")
    async def gm_ghep_lai(request: Request):
        """Ghép lại toàn bộ hoá đơn Gmail vào kho Tải Hoá Đơn TCT (vd sau khi vừa tra cứu thêm kỳ mới,
        hay vừa thêm doanh nghiệp). Chỉ cần mở khoá Quét Gmail."""
        if (c := _chan(APP_GMAIL)):
            return c
        d = await _than(request)
        try:
            ids = {r["id"] for r in chon_gmail(d)}
        except ValueError as e:
            return _loi(e)
        with _KHOA_GHI:
            kho = _doc_json(_gmail_kho_path(), {})
            dem = {"ghep": 0, "moi": 0}
            for ban in kho.values():
                if ban.get("id") not in ids:
                    continue
                if ban.get("lien_ket") and tuyet_doi(ban.get("pdf")) is None:
                    continue
                ket = lien_ket_gmail(ban)
                if ket:
                    dem[ket] += 1
            _ghi_json(_gmail_kho_path(), kho)
        return {"ok": True, **dem}

    @app.post(A + "/gmail/xuat")
    async def gm_xuat(request: Request):
        if (c := _chan(APP_GMAIL)):
            return c
        try:
            ds = chon_gmail(await _than(request))
        except ValueError as e:
            return _loi(e)
        if not ds:
            return _loi(ValueError("Chưa quét được hoá đơn nào."))
        return _tai_ve(xuat_gmail_excel(ds), f"HoaDonGmail_{dt.date.today():%Y%m%d}.xlsx", XLSX)

    @app.post(A + "/gmail/zip")
    async def gm_zip(request: Request):
        if (c := _chan(APP_GMAIL)):
            return c
        try:
            ds = chon_gmail(await _than(request))
        except ValueError as e:
            return _loi(e)
        buf = io.BytesIO()
        n = 0
        manifest = []
        kho_tct = {}
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for index, r in enumerate(ds, 1):
                for k in ("xml", "pdf", "trong_dong"):
                    f = gmail_trong_dong(r, kho_tct) if k == "trong_dong" else tuyet_doi(r.get(k))
                    if f:
                        z.write(f, f"{index:04d}/{k}/{f.name}")
                        n += 1
                    manifest.append({"id": r["id"], "loai": k, "da_co": bool(f)})
            z.writestr("doi-soat-file.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        if not n:
            return _loi(ValueError("Chưa có file nào."))
        return _tai_ve(buf.getvalue(), f"HoaDonGmail_{dt.date.today():%Y%m%d}.zip", "application/zip")

    @app.post(A + "/gmail/pdf-gop")
    async def gm_gop(request: Request):
        if (c := _chan(APP_GMAIL)):
            return c
        try:
            ds = chon_gmail(await _than(request))
            files = [tuyet_doi(r.get("pdf")) for r in ds]
            missing = sum(not f for f in files)
            if missing:
                raise ValueError(f"Có {missing}/{len(ds)} hoá đơn chưa có PDF. Tải chứng từ trước hoặc lọc 'Có PDF gốc'.")
            if not files:
                raise ValueError("Chưa chọn hoá đơn nào.")
            import tempfile
            with tempfile.TemporaryDirectory() as td:
                out = str(Path(td) / "gop.pdf")
                n = thd("exporter").merge_pdfs([str(f) for f in files], out)
                return _tai_ve(Path(out).read_bytes(), f"Gmail_IN_GOP_{n}HD.pdf", "application/pdf")
        except Exception as e:
            return _loi(e)

    @app.post(A + "/gmail/tai")
    async def gm_tai(request: Request):
        if (c := _chan(APP_GMAIL)):
            return c
        try:
            ds = chon_gmail(await _than(request))
            if not ds:
                raise ValueError("Chưa chọn hoá đơn nào.")
            v = chay_viec("gmail", f"Tải chứng từ {len(ds)} hoá đơn Gmail", lambda v: tai_gmail_da_chon(v, ds))
            return {"ok": True, "viec": v.id}
        except Exception as e:
            return _loi(e)

    @app.post(A + "/gmail/ma-tra-cuu")
    async def gm_ma(request: Request):
        if (c := _chan(APP_GMAIL)):
            return c
        d = await _than(request)
        code = str(d.get("code") or "").strip()
        if not re.fullmatch(r"[A-Za-z0-9._-]{0,200}", code):
            return _loi(ValueError("Mã tra cứu không hợp lệ."))
        with _KHOA_GHI:
            kho = _doc_json(_gmail_kho_path(), {})
            r = kho.get(str(d.get("id")))
            if r is None:
                return _loi(ValueError("Không tìm thấy hoá đơn."), 404)
            r["ma_tra_cuu"] = code
            if r.get("nha_cung_cap") == "bkav" and code:
                from urllib.parse import quote
                r["link"] = "https://tchd.ehoadon.vn/TCHD?MTC=" + quote(code, safe="")
            _ghi_json(_gmail_kho_path(), kho)
        return {"ok": True}

    # ---------------- file tĩnh (đặt CUỐI để không nuốt đường /api)
    @app.get(GOC + "/{ten}")
    def hd_tinh(ten: str):
        if ten not in WEB_FILES:
            return JSONResponse({"error": "not found"}, status_code=404)
        if ly_do_chan(APP_TAI) and ly_do_chan(APP_GMAIL):
            return JSONResponse({"error": "locked"}, status_code=403)
        return FileResponse(str(WEB_DIR / ten), headers={"Cache-Control": "no-cache"})
