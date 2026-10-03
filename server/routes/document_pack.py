"""Explicit document selection and bounded staff spreadsheet intake."""
import base64
import datetime as dt
import io
import re
import tempfile
import unicodedata
import zipfile
from pathlib import Path

FIELDS = ['ho_ten', 'ngay_sinh', 'cccd', 'dia_chi', 'chuc_vu', 'luong', 'ngay_vao', 'ma_bhxh']
HEADERS = ['Họ tên', 'Ngày sinh', 'CCCD', 'Địa chỉ', 'Chức vụ', 'Lương', 'Ngày vào làm', 'Mã BHXH']
CATALOG = []
def item(id, tab, title, filename, source, index=0):
    CATALOG.append(dict(id=id, tab=tab, title=title, filename=filename, source=source, index=index))

for n, title in enumerate(['Phiếu nhập kho', 'Đề nghị thanh toán', 'Phiếu chi / Ủy nhiệm chi', 'Biên bản giao nhận', 'Đối chiếu công nợ phải trả']):
    item('purchase_'+str(n), 'invoice', title, '', 'purchase', n)
for n, title in enumerate(['Phiếu xuất kho', 'Phiếu thu', 'Bàn giao / nghiệm thu', 'Đối chiếu công nợ phải thu']):
    item('sold_'+str(n), 'invoice', title, '', 'sold', n)
for n, title in enumerate(['Hợp đồng nguyên tắc / mua bán', 'Nghiệm thu / thanh lý hợp đồng']):
    item('contract_'+str(n), 'invoice', title, '', 'contract', n)
for n, title in enumerate(['Kiểm kê kho', 'Phiếu giao nhận chung', 'Đơn đặt gia công', 'Nghiệm thu gia công', 'Hợp đồng gia công']):
    item('process_'+str(n), 'invoice', title, '', 'process', n)
for n, title in enumerate(['Hợp đồng lao động', 'Quyết định tiếp nhận', 'Quyết định lương', 'Cam kết bảo mật / bàn giao', 'Khung thông tin TK1-TS / BHXH']):
    item('staff_'+str(n), 'labor', title, '', 'staff', n)
item('payroll', 'labor', 'Bảng lương (Excel)', 'Bang_Luong.xlsx', 'sheet')
item('attendance', 'labor', 'Bảng chấm công (Excel)', 'Bang_Cham_Cong.xlsx', 'sheet')
item('insurance', 'labor', 'Danh sách BHXH D02-LT — khung dữ liệu (Excel)', 'Khung_D02_LT.xlsx', 'sheet')
item('vat', 'compliance', 'Bảng kê hóa đơn / VAT (Excel)', 'Bang_Ke_VAT.xlsx', 'sheet')
item('tax', 'compliance', 'Checklist hồ sơ thuế', 'Checklist_Thue.doc', 'tax')
for n, title in enumerate(['Checklist lưu trữ và PCCC', 'Biên bản tự kiểm tra PCCC', 'Phân công PCCC / lưu trữ', 'Lịch kiểm tra định kỳ']):
    item('safety_'+str(n), 'compliance', title, '', 'safety', n)

def norm(value):
    value = unicodedata.normalize('NFD', str(value or '').lower().replace('đ', 'd'))
    return re.sub(r'[^a-z0-9]', '', ''.join(c for c in value if not unicodedata.combining(c)))

def workbook(headers, rows, title):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = title
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for row in ws:
        for c in row:
            # User-provided values must never become spreadsheet formulas.
            if isinstance(c.value, str): c.data_type = 's'
            c.font = Font(name='Arial', size=11)
            c.alignment = Alignment(vertical='top', wrap_text=True)
    for c in ws[1]:
        c.fill = PatternFill('solid', fgColor='6046D8'); c.font = Font(name='Arial', bold=True, color='FFFFFF')
    for i, h in enumerate(headers, 1): ws.column_dimensions[get_column_letter(i)].width = 25 if i > 1 else 32
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    ws.print_title_rows = '1:1'; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    return wb

def excel_bytes(wb):
    b = io.BytesIO(); wb.save(b); return b.getvalue()

def staff_template():
    wb = workbook(HEADERS, [], 'Nhan_su'); ws = wb.active
    for row in ws.iter_rows(min_row=2, max_row=501, max_col=8):
        for c in row: c.number_format = '@'
    guide = wb.create_sheet('Huong_dan')
    for text in ['Nhập dữ liệu ở sheet Nhan_su, dòng 2 trở đi. Tối đa 500 người.',
                 'Có thể đổi thứ tự cột; hệ thống khớp theo tên cột. Họ tên bắt buộc.',
                 'CCCD và Mã BHXH giữ kiểu Text để không mất số 0 đầu.',
                 'Ngày: dd/mm/yyyy hoặc yyyy-mm-dd. Lương: số tiền VND không âm.',
                 'App sẽ xem trước và báo lỗi theo dòng trước khi bạn áp dụng vào bảng.']:
        guide.append([text])
    guide.column_dimensions['A'].width = 110
    return excel_bytes(wb)

def parse_staff(data):
    from openpyxl import load_workbook
    if len(data) > 5*1024*1024: raise ValueError('Excel tối đa 5 MB.')
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if len(z.infolist()) > 1000 or sum(x.file_size for x in z.infolist()) > 30*1024*1024:
            raise ValueError('Excel vượt giới hạn giải nén.')
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        ws = wb['Nhan_su'] if 'Nhan_su' in wb.sheetnames else wb.active
        aliases = {norm(h): f for h, f in zip(HEADERS, FIELDS)}
        aliases.update({norm(f): f for f in FIELDS})
        aliases.update({'hovaten':'ho_ten','hotennhanvien':'ho_ten','socccd':'cccd','chucdanh':'chuc_vu','luongcoban':'luong','masobhxh':'ma_bhxh','ngayvaolam':'ngay_vao'})
        iterator = ws.iter_rows(); header = next(iterator, ())
        mapping = {i:aliases[norm(c.value)] for i,c in enumerate(header) if norm(c.value) in aliases}
        if 'ho_ten' not in mapping.values(): raise ValueError('Không tìm thấy cột Họ tên ở dòng đầu.')
        if len(set(mapping.values())) != len(mapping): raise ValueError('Có cột nhân sự trùng tên; hãy giữ một cột cho mỗi thông tin.')
        rows, errors, seen = [], [], set()
        for line, cells in enumerate(iterator, 2):
            if line > 501: raise ValueError('Tối đa 500 dòng nhân sự; tách file rồi nhập tiếp.')
            if not any(c.value is not None for c in cells): continue
            row = {f:'' for f in FIELDS}; issues=[]
            for i,f in mapping.items():
                if i >= len(cells): continue
                c=cells[i]; value=c.value
                if c.data_type == 'f': issues.append('Không nhận công thức ở '+HEADERS[FIELDS.index(f)]); continue
                if isinstance(value, (dt.date,dt.datetime)): value=value.strftime('%d/%m/%Y')
                if isinstance(value, float) and value.is_integer(): value=int(value)
                value=str(value if value is not None else '').strip()
                if len(value)>300: issues.append('Ô quá dài'); continue
                if f in ('ngay_sinh','ngay_vao') and value:
                    date=None
                    for fmt in ('%d/%m/%Y','%Y-%m-%d'):
                        try: date=dt.datetime.strptime(value,fmt); break
                        except ValueError: pass
                    if date: value=date.strftime('%d/%m/%Y')
                    else: issues.append('Ngày không hợp lệ: '+value)
                if f in ('cccd','ma_bhxh') and value:
                    if c.data_type == 'n': issues.append(HEADERS[FIELDS.index(f)]+' phải là Text để giữ số 0 đầu')
                    if not value.isdigit(): issues.append(HEADERS[FIELDS.index(f)]+' chỉ chứa chữ số')
                if f=='luong' and value:
                    clean=re.sub(r'[.,\s]', '', value)
                    if not clean.isdigit(): issues.append('Lương phải là số tiền không âm')
                    else: value=clean
                row[f]=value
            if not row['ho_ten']: issues.append('Thiếu họ tên')
            key=row['cccd'] or row['ma_bhxh'] or norm(row['ho_ten'])+'|'+row['ngay_sinh']
            if key in seen: issues.append('Trùng người trong file')
            seen.add(key)
            if issues: errors.append({'line':line,'message':'; '.join(issues)})
            else: rows.append(row)
        return {'rows':rows,'errors':errors,'columns':[HEADERS[FIELDS.index(f)] for f in mapping.values()]}
    finally: wb.close()

def selection(d):
    tab=d.get('document_tab'); ids=d.get('templates')
    if tab not in ('invoice','labor','compliance') or not isinstance(ids,list) or not ids:
        raise ValueError('Chọn tab và ít nhất một mẫu chứng từ.')
    selected=[x for x in CATALOG if x['id'] in ids]
    if len(selected)!=len(set(ids)) or any(x['tab']!=tab for x in selected):
        raise ValueError('Mẫu không thuộc mục đang tải. Hãy tải riêng từng mục.')
    return selected

def documents(hd, d, p, s):
    chosen=selection(d); cty=hd._cty(d['mst']); style=d.get('kieu_mau','ke_toan')
    if not isinstance(d.get('nhan_su', []), list) or len(d.get('nhan_su', [])) > 500:
        raise ValueError('Tối đa 500 người lao động mỗi gói.')
    people=hd._nhan_su_sach(d.get('nhan_su')); docs=[]; sheets=[]
    if d['document_tab']=='labor' and not people: raise ValueError('Nhập ít nhất một người lao động.')
    sizes={'purchase':len(p),'sold':len(s),'staff':len(people)}
    if sum(sizes.get(m['source'],1) for m in chosen)>1500:
        raise ValueError('Tối đa 1.500 chứng từ mỗi gói; hãy chia nhỏ lựa chọn.')
    for m in chosen:
        src=m['source']; n=m['index']; id=m['id']
        if src in ('purchase','sold'):
            rows=p if src=='purchase' else s
            if not rows: raise ValueError('Mẫu '+m['title']+' cần hóa đơn '+('mua vào' if src=='purchase' else 'bán ra')+'.')
            for i,inv in enumerate(rows,1):
                fn,data=hd._invoice_docs(cty,inv,src,{'mua_vao','ban_ra'},style)[n]
                folder=f"HoaDon/{i:03d}_{hd._slug(hd._hoa_don_id(inv))}_{hd._slug(hd._ben(inv)['ten'],40)}"
                docs.append((folder+'/'+fn,data))
        elif src=='staff':
            docs += [('LaoDong/'+fn,data) for fn,data in hd._mau_lao_dong(cty,style,people) if Path(fn).name.startswith(f'{n+1:02d}_')]
        elif src in ('contract','process','safety'):
            factory={'contract':hd._mau_hop_dong,'process':hd._mau_gia_cong,'safety':hd._mau_pccc_luu_tru}[src]
            fn,data=factory(cty,style)[n]; docs.append((fn,data))
        elif src=='tax':
            body='<h1>Checklist hồ sơ thuế</h1><table><tr><th>Hồ sơ</th><th>Đã có</th><th>Ghi chú</th></tr>'+''.join('<tr><td>'+t+'</td><td></td><td></td></tr>' for t in ['Hóa đơn XML/PDF','Bảng kê VAT','Hợp đồng và giao nhận','Chứng từ thanh toán','Giải trình chênh lệch'])+'</table>'
            docs.append((m['filename'],hd._doc_bytes(m['title'],body,style)))
        else:
            if id=='vat':
                if not p+s: raise ValueError('Chọn hóa đơn để tạo bảng kê VAT.')
                headers=['Ngày','Số','Ký hiệu','Đối tác','MST','Tiền hàng','Thuế','Tổng']
                rows=[[hd._ngay_vn(hd.ngay_hd(r)),str(r.get('shdon','')),r.get('khhdon',''),hd._ben(r)['ten'],hd._ben(r)['mst'],r.get('tien_hang',r.get('tgtcthue',0)),r.get('thue',r.get('tgtthue',0)),r.get('tong',r.get('tgtttbso',0))] for r in p+s]
            elif id=='attendance':
                headers=['Họ tên']+[str(i) for i in range(1,32)]+['Tổng công']; rows=[[r['ho_ten']]+['']*32 for r in people]
            else:
                fields=['ho_ten','chuc_vu','luong'] if id=='payroll' else FIELDS
                headers=[HEADERS[FIELDS.index(f)] for f in fields]; rows=[[r[f] for f in fields] for r in people]
            sheets.append((m['filename'],excel_bytes(workbook(headers,rows,id))))
            # Excel selections are also included in the single printable packet.
            body='<h1>'+hd._html(m['title'])+'</h1><table><thead><tr>'+''.join('<th>'+hd._html(h)+'</th>' for h in headers)+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+hd._html(str(v))+'</td>' for v in row)+'</tr>' for row in rows)+'</tbody></table>'
            docs.append((m['filename']+'.print',hd._doc_bytes(m['title'],body,style)))
    if len(docs)>1500: raise ValueError('Tối đa 1.500 chứng từ mỗi gói; hãy chia nhỏ lựa chọn.')
    return docs,sheets

def combined(hd, docs, style):
    bodies=[]
    for name,data in docs:
        text=data.decode('utf-8'); body=re.search(r'<body>(.*)</body>',text,re.S).group(1)
        bodies.append('<section class="document'+(' spreadsheet' if name.endswith('.print') else '')+'">'+body+'</section>')
    html=hd._doc_frame('Hồ sơ đã chọn — in gộp',''.join(bodies),style)
    return html.replace('</style>', '@page spreadsheet{size:A4 landscape;margin:12mm}.spreadsheet{page:spreadsheet;font-size:9pt}.spreadsheet td,.spreadsheet th{padding:3px;overflow-wrap:anywhere}.spreadsheet table{table-layout:fixed}.document{break-before:page}.document:first-child{break-before:auto}.page{border:0;box-shadow:none;padding:0}thead{display:table-header-group}tr{break-inside:avoid}@media print{body{background:white}.print-button{display:none}}\n</style>').replace('<body>','<body><button class="print-button" onclick="window.print()">In toàn bộ hồ sơ đã chọn</button>')

def export(hd,d,p,s):
    docs,sheets=documents(hd,d,p,s); html=combined(hd,docs,d.get('kieu_mau','ke_toan'))
    buf=io.BytesIO()
    with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in docs+sheets:
            if not name.endswith('.print'): z.writestr(name,data)
        z.writestr('00_IN_GOP.html',html.encode('utf-8'))
        browser=hd.trinh_duyet_in_pdf()
        if browser:
            with tempfile.TemporaryDirectory(prefix='minmin_documents_') as tmp:
                src=Path(tmp)/'print.html'; dst=Path(tmp)/'print.pdf'; src.write_text(html,encoding='utf-8')
                hd.thd('exporter').html_to_pdf(str(src),str(dst),browser,timeout=120)
                z.writestr('00_IN_GOP.pdf',dst.read_bytes())
    return buf.getvalue()

def register(app,hd,choose):
    from fastapi import Request
    A=hd.GOC+'/api/chung-tu'
    @app.get(A+'/mau')
    def catalog():
        if (c:=hd._chan_chung_tu()): return c
        return {'templates':CATALOG}
    @app.get(A+'/nhan-su-mau')
    def template():
        if (c:=hd._chan_chung_tu()): return c
        return hd._tai_ve(staff_template(),'Mau_Nhan_Su.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    @app.post(A+'/nhan-su-nhap')
    async def intake(request:Request):
        if (c:=hd._chan_chung_tu()): return c
        try:
            buf=bytearray()
            async for chunk in request.stream():
                buf.extend(chunk)
                if len(buf)>7*1024*1024: raise ValueError('Excel tối đa 5 MB.')
            import json
            d=json.loads(buf); return parse_staff(base64.b64decode(d.get('data',''),validate=True))
        except Exception as e: return hd._loi(e)
    @app.post(A+'/xem-mau')
    async def preview(request:Request):
        if (c:=hd._chan_chung_tu()): return c
        try:
            d=await hd._than(request); d['mst']=hd.chuan_mst(d.get('mst'))
            (p,s),_,_=choose(d['mst'],d)
            docs,_=documents(hd,d,p[:1],s[:1]); return {'html':combined(hd,docs[:1],d.get('kieu_mau','ke_toan'))}
        except Exception as e: return hd._loi(e)
