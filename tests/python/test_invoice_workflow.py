"""Selected exports, encrypted captcha config, lookup codes; no paid/network calls."""
from _paths import ROOT, SERVER
import io
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfWriter, PdfReader
import routes.hoa_don as hd

A = '/apps/tai-hoa-don/api'

class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        for p in [patch.object(hd,'goc_du_lieu',return_value=self.root),patch.object(hd,'ly_do_chan',return_value=None)]:
            p.start(); self.addCleanup(p.stop)
        app=FastAPI(); hd.register(app); self.c=TestClient(app)
        self.addCleanup(self.c.close)
        w=PdfWriter(); w.add_blank_page(width=200,height=200)
        with (self.root/'a.pdf').open('wb') as f: w.write(f)
        (self.root/'a.xml').write_text('<HDon/>')
        hd._ghi_json(hd._gmail_kho_path(),{'a':{'id':'a','so':'10','pdf':'a.pdf','xml':'a.xml','ma_tra_cuu':'SAMPLE123','nha_cung_cap':'bkav'},'b':{'id':'b','so':'20'}})

    def test_selected_exports_and_empty_never_exports_all(self):
        r=self.c.post(A+'/gmail/xuat',json={'ids':['b']})
        self.assertEqual(r.status_code,200)
        ws=load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.max_row,2); self.assertEqual(ws.cell(2,2).value,'20')
        r=self.c.post(A+'/gmail/zip',json={'ids':['a']})
        with zipfile.ZipFile(io.BytesIO(r.content)) as z:
            self.assertEqual(len(z.namelist()),3)
        for endpoint in ('xuat','zip','pdf-gop','tai','ghep-lai'):
            self.assertEqual(self.c.post(A+'/gmail/'+endpoint,json={'ids':[]}).status_code,400)
            self.assertEqual(self.c.post(A+'/gmail/'+endpoint,json={'ids':['missing']}).status_code,400)
        self.assertEqual(self.c.post(A+'/xuat',json={'mst':'1234567890','loai':'purchase','keys':[]}).status_code,400)
        self.assertEqual(self.c.post(A+'/tai',json={'mst':'1234567890','loai':'purchase','keys':[]}).status_code,400)

    def test_merge_refuses_missing_pdf(self):
        r=self.c.post(A+'/gmail/pdf-gop',json={'ids':['a']})
        self.assertEqual(r.status_code,200,r.text[:100] if r.status_code!=200 else '')
        self.assertEqual(len(PdfReader(io.BytesIO(r.content)).pages),1)
        self.assertEqual(self.c.post(A+'/gmail/pdf-gop',json={'ids':['a','b']}).status_code,400)

    def test_document_types_and_linked_tax_rendering(self):
        hd.ghi_cau_hinh({'doanh_nghiep':{'1234567890':{'ten':'Fixture'}}})
        hd.ghi_kho('1234567890','purchase',{'invoice':{'_pdf_path':'a.pdf'}})
        kho=hd._doc_json(hd._gmail_kho_path(),{})
        kho['a']['lien_ket']={'mst':'1234567890','key':'invoice'}
        hd._ghi_json(hd._gmail_kho_path(),kho)
        rows={r['id']:r for r in self.c.get(A+'/gmail/ket-qua').json()['ds']}
        self.assertTrue(rows['a']['co_goc']); self.assertTrue(rows['a']['co_trong_dong'])
        self.assertFalse(rows['b']['co_goc']); self.assertFalse(rows['b']['co_trong_dong'])
        r=self.c.get(A+'/gmail/tep/a/trong_dong')
        self.assertEqual(r.status_code,200); self.assertTrue(r.content.startswith(b'%PDF'))
        self.assertIn('application/pdf',r.headers['content-type'])
        with zipfile.ZipFile(io.BytesIO(self.c.post(A+'/gmail/zip',json={'ids':['a']}).content)) as z:
            self.assertEqual(len(z.namelist()),4)
            self.assertTrue(any('/trong_dong/' in n for n in z.namelist()))
        # An unconfigured company or deleted file must not be advertised as downloadable.
        hd.ghi_cau_hinh({'doanh_nghiep':{}})
        self.assertEqual(self.c.get(A+'/gmail/tep/a/trong_dong').status_code,404)
        self.assertFalse(hd.tom_tat({'_pdf_path':'missing.pdf'})['co_trong_dong'])

    def test_key_encryption_mask_and_provider_switch(self):
        r=self.c.post(A+'/captcha-api',json={'provider':'2captcha','api_key':'fake-secret','enabled':True})
        self.assertEqual(r.status_code,200); self.assertNotIn('fake-secret',r.text)
        self.assertNotIn('fake-secret',hd._cau_hinh_path().read_text())
        self.assertTrue(r.json()['keys']['2captcha'])
        self.assertEqual(hd.tao_solver().services,[('2captcha','fake-secret')])
        self.c.post(A+'/captcha-api',json={'provider':'2captcha','api_key':'','enabled':True})
        self.assertEqual(hd.tao_solver().services,[('2captcha','fake-secret')])
        self.assertEqual(self.c.post(A+'/captcha-api',json={'provider':'omocaptcha','enabled':True}).status_code,400)
        self.c.post(A+'/captcha-api',json={'provider':'2captcha','xoa_key':True,'enabled':False})
        self.assertIsNone(hd.tao_solver())

    def test_bkav_code_url_and_persistence(self):
        r=self.c.post(A+'/gmail/ma-tra-cuu',json={'id':'a','code':'HELLO123'})
        self.assertEqual(r.status_code,200)
        row=hd.chon_gmail({'ids':['a']})[0]
        self.assertIn('?MTC=HELLO123',row['link'])
        self.assertEqual(self.c.post(A+'/gmail/ma-tra-cuu',json={'id':'a','code':'<script>'}).status_code,400)
        xml=b'<HDon><MSTTCGP>0101360697</MSTTCGP></HDon>'
        info=hd.thd('providers').manual_download_info(xml,{'_bkav_mtc':'HELLO123'})
        self.assertEqual(info['code'],'HELLO123'); self.assertIn('MTC=HELLO123',info['url'])
        for provider in ('bkav','easyinvoice'):
            mod=hd.thd('providers.'+provider)
            self.assertEqual(mod.find_code('<HDon><TTin><TTruong>MaTraCuu</TTruong><DLieu>OLD123</DLieu></TTin></HDon>',{'_ma_tra_cuu':'NEW456'}),'NEW456')

    def test_solver_wired_into_retry(self):
        with patch.object(hd,'tao_solver',return_value='mock-solver'), patch.object(hd.thd('providers'),'fetch_original_pdf',return_value=('test',(self.root/'a.pdf').read_bytes())) as fetch:
            r=hd.chon_gmail({'ids':['a']})[0]; r['pdf']=None
            hd.tai_gmail_da_chon(hd.Viec('gmail','test'),[r])
            self.assertEqual(fetch.call_args.kwargs['solver'],'mock-solver')

    def test_captcha_task_contracts_without_network(self):
        solver=hd.thd('captcha_solver')
        for provider,field,task in [('2captcha','body','RecaptchaV2TaskProxyless'),('omocaptcha','imageBase64','RecaptchaV2TokenTask')]:
            with patch.object(solver,'_solve',return_value='ABC') as solve:
                self.assertEqual(solver.solve_image(provider,'test',b'img'),'ABC')
                self.assertIn(field,solve.call_args.args[2])
                solver.solve_recaptcha_v2(provider,'test','site-key','https://example.com/invoice')
                self.assertEqual(solve.call_args.args[2],{'type':task,'websiteURL':'https://example.com/invoice','websiteKey':'site-key'})
        with patch.object(hd,'ly_do_chan',return_value='locked'):
            self.assertEqual(self.c.get(A+'/captcha-api').status_code,403)
            self.assertEqual(self.c.post(A+'/gmail/pdf-gop',json={'ids':['a']}).status_code,403)

if __name__=='__main__': unittest.main()
