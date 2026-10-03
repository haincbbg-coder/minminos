from _paths import ROOT, SERVER
import base64
import io
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from routes import hoa_don as hd, document_pack as pack

INV={'nbmst':'001203013937','nbten':'PHƯƠNG HUY','khmshdon':2,'khhdon':'C26MYY','shdon':168,'tdlap':'2026-10-02','tgtttbso':30000000,'_loai':'purchase'}
PERSON=dict(zip(pack.FIELDS,['Nguyễn A','01/01/1990','001090123456','Hà Nội','Kế toán','8000000','01/01/2026','0012345678']))

class DocumentPackTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        for p in [patch.object(hd,'goc_du_lieu',return_value=Path(self.tmp.name)),patch.object(hd,'ly_do_chan',return_value=None),patch.object(hd,'trinh_duyet_in_pdf',return_value='')]:
            p.start();self.addCleanup(p.stop)
        hd.ghi_cau_hinh({'doanh_nghiep':{'2401063265':{'ten':'HALIMI'}}})
        hd.ghi_kho('2401063265','purchase',{hd.khoa_hd(INV):INV})
        app=FastAPI();hd.register(app);self.c=TestClient(app);self.addCleanup(self.c.close)

    def data(self,tab,ids):
        return dict(mst='2401063265',loai='purchase',keys=[hd.khoa_hd(INV)],document_tab=tab,templates=ids,nhan_su=[PERSON])

    def export(self,d):
        r=self.c.post('/apps/tai-hoa-don/api/bo-chung-tu',json=d)
        self.assertEqual(r.status_code,200,r.text if r.status_code!=200 else '')
        return zipfile.ZipFile(io.BytesIO(r.content))

    def test_three_invoice_templates_only_and_combined_contains_all_three(self):
        z=self.export(self.data('invoice',['purchase_0','purchase_1','purchase_3']))
        self.assertEqual(len(z.namelist()),4)
        self.assertEqual(sum(n.endswith('.doc') for n in z.namelist()),3)
        html=z.read('00_IN_GOP.html').decode()
        self.assertEqual(html.count('<section class="document">'),3)
        self.assertIn('PHƯƠNG HUY',html);self.assertNotIn('Bảng lương',html)

    def test_compliance_has_no_hidden_labor_invoice_or_checklist57(self):
        z=self.export(self.data('compliance',['safety_1']))
        self.assertEqual(set(z.namelist()),{'BienBan_TuKiemTra_PCCC.doc','00_IN_GOP.html'})

    def test_labor_only_selected_template_and_person(self):
        d=self.data('labor',['staff_2']);d['keys']=[]
        z=self.export(d);self.assertEqual(len(z.namelist()),2)
        self.assertTrue(any('03_QuyetDinhLuongPhuCap.doc' in n for n in z.namelist()))
        self.assertIn('Nguyễn A',z.read('00_IN_GOP.html').decode())

    def test_vat_workbook_has_only_selected_sheet(self):
        z=self.export(self.data('compliance',['vat']))
        wb=load_workbook(io.BytesIO(z.read('Bang_Ke_VAT.xlsx')))
        self.assertEqual(wb.sheetnames,['vat']);self.assertEqual(wb.active.max_row,2)
        self.assertNotIn('Bang_Luong.xlsx',z.namelist())

    def test_empty_and_cross_tab_selection_rejected(self):
        for tab,ids in [('compliance',['staff_0']),('invoice',[]),('labor',['unknown'])]:
            self.assertEqual(self.c.post('/apps/tai-hoa-don/api/bo-chung-tu',json=self.data(tab,ids)).status_code,400)

    def test_staff_template_reordered_headers_preserves_leading_zero(self):
        wb=load_workbook(io.BytesIO(pack.staff_template()));ws=wb.active
        for i,(h,f) in enumerate(reversed(list(zip(pack.HEADERS,pack.FIELDS))),1):
            ws.cell(1,i,h);ws.cell(2,i,PERSON[f])
        result=pack.parse_staff(pack.excel_bytes(wb))
        self.assertEqual(result['errors'],[]);self.assertEqual(result['rows'],[PERSON])

    def test_import_reports_invalid_dates_numeric_id_formula_and_duplicate(self):
        wb=pack.workbook(pack.HEADERS,[[PERSON[f] for f in pack.FIELDS]],'Nhan_su');ws=wb.active
        ws['B2']='31/02/2026';ws['C2']=1090123456;ws['F2']='=1+2'
        ws.append([PERSON[f] for f in pack.FIELDS]);ws.append([PERSON[f] for f in pack.FIELDS])
        result=pack.parse_staff(pack.excel_bytes(wb))
        self.assertEqual([e['line'] for e in result['errors']],[2,4])
        self.assertIn('Text',result['errors'][0]['message']);self.assertIn('công thức',result['errors'][0]['message'])

    def test_preview_uses_selected_invoice_and_escapes_content(self):
        d=self.data('invoice',['purchase_1'])
        r=self.c.post('/apps/tai-hoa-don/api/chung-tu/xem-mau',json=d)
        self.assertEqual(r.status_code,200,r.text);self.assertIn('PHƯƠNG HUY',r.json()['html'])
        self.assertIn('168',r.json()['html'])

    def test_unknown_or_empty_invoice_scope_does_not_export_all(self):
        d=self.data('invoice',['purchase_0']);d['keys']=[]
        self.assertEqual(self.c.post('/apps/tai-hoa-don/api/bo-chung-tu',json=d).status_code,400)

    def test_excel_formula_injection_is_text(self):
        wb=pack.workbook(['Name'],[['=HYPERLINK("https://example.com")']],'test')
        cell=load_workbook(io.BytesIO(pack.excel_bytes(wb))).active['A2']
        self.assertEqual(cell.data_type,'s')

if __name__=='__main__':unittest.main()
